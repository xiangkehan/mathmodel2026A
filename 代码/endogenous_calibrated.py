"""内生收缩模型（终版）：物理结构化闭合 A/B/C 比选 + 附件 2 标定 + 全套验证。

闭合 C（最终采用，全湿度驱动、一项一机制）：
  R(t) = R0·sqrt((1+Ū)/(1+C0))          失水收缩（水分守恒 + 等密度假设）
        · [1 − ε_c·g_c(Ū)]               毛细-膨压塌陷：g_c(x)=(1−x)^a·x^b（归一化），
                                          x=失水率；先升后降=塌陷发展→成孔回弹
        · [1 + ε_p·g_p(C_s)]             结皮-成孔停滞：g_p=clip((C_glass−C_s)/Δ,0,1)^k_p，
                                          C_s 穿越 C_glass=0.21（文献值，非拟合）物理触发
标定参数（5 个，拟合附件 2，逐项标注来源）：ε_c, a, b, ε_p, k_p。

用法：cd 代码 && python endogenous_calibrated.py
产物：../数据/endogenous_calibrated.csv、../数据/endogenous_coupled.csv（数值产物）；
      ../08-临时/内生收缩_标定验证.png（预览图，非论文交付图，论文用 fig7_内生收缩贴合）。
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import least_squares
from scipy.interpolate import PchipInterpolator

HERE = Path(__file__).resolve().parent          # 代码
A_DIR = HERE.parent                             # 项目根（A/）
sys.path.insert(0, str(HERE))
DATA = A_DIR / "数据"
SCRATCH = A_DIR / "08-临时"                     # 预览图等非交付产物

from solver_q1 import R0, C0
from solver_q4 import PROPS4, D_of4, surface_m4, coupled_step4, solve_coupled4
from solver_q23 import TH
from run_q23 import read_env, make_env

C_GLASS = 0.21                        # 文献/Gordon–Taylor（非拟合）
RHO_SK, RHO_W = 1468.0, 1000.0        # 骨架密度（文献值，非拟合）

# ---------- 数据 ----------
wb = openpyxl.load_workbook(A_DIR / "01-题目" / "原始文件" / "附件2.xlsx", read_only=True)
rows = [r for r in wb[wb.sheetnames[0]].iter_rows(min_row=2, values_only=True) if r[0] is not None]
wb.close()
t_att = np.array([r[0] for r in rows], float)
R_att = np.array([r[1] for r in rows], float) * 0.01
d = np.load(DATA / "q4_main_steps.npz")
um = np.interp(t_att, d["t"], d["um"])
Cs = np.interp(t_att, d["t"], d["Cs"])
th = t_att / 3600.0
CSMIN = float(Cs.min())               # 全局归一化常数（末期 C_s，物理量）
R_ideal = R0 * np.sqrt((1 + um) / (1 + C0))
m_data = R_att / R_ideal


# ---------- 三个候选闭合 ----------
def closureA(p, t, Cs_):
    eps_c, tau_c, eps_p, t_p, tau_p = p
    Ri = R0 * np.sqrt((1 + np.interp(t, t_att, um)) / (1 + C0)) if t is not t_att else R_ideal
    phi_c = 1.0 - np.exp(-t / tau_c)
    phi_p = 1.0 / (1.0 + np.exp(-(t - t_p) / tau_p))
    return Ri * (1.0 - eps_c * phi_c) * (1.0 + eps_p * phi_p)


def closureB(p, t, Cs_):
    eps_c, tau_c, tau_d, eps_p, k_p = p
    phi_c = (1.0 - np.exp(-t / tau_c)) * np.exp(-t / tau_d)
    phi_p = np.clip((C_GLASS - Cs_) / (C_GLASS - CSMIN), 0.0, 1.0) ** k_p
    return R_ideal * (1.0 - eps_c * phi_c) * (1.0 + eps_p * phi_p)


def closureC(p, um_, Cs_):
    eps_c, a, b, eps_p, k_p = p
    Ri = R0 * np.sqrt((1 + um_) / (1 + C0))
    x = np.clip((C0 - um_) / C0, 0, 1)
    gc = (1 - x) ** a * x ** b
    gc = gc / ((1 - b / (a + b)) ** a * (b / (a + b)) ** b)     # 峰值归一
    gp = np.clip((C_GLASS - Cs_) / (C_GLASS - CSMIN), 0.0, 1.0) ** k_p
    return Ri * (1 - eps_c * gc) * (1 + eps_p * gp)


def termsC(p, um_, Cs_):
    eps_c, a, b, eps_p, k_p = p
    x = np.clip((C0 - um_) / C0, 0, 1)
    gc = (1 - x) ** a * x ** b
    gc = gc / ((1 - b / (a + b)) ** a * (b / (a + b)) ** b)
    gp = np.clip((C_GLASS - Cs_) / (C_GLASS - CSMIN), 0.0, 1.0) ** k_p
    return -eps_c * gc, eps_p * gp


def fit(fn, args, p0, lb, ub, target=None):
    if target is None:
        target = R_att
    s = least_squares(lambda p: fn(p, *args) - target, p0, bounds=(lb, ub),
                      xtol=1e-15, ftol=1e-15, gtol=1e-15, max_nfev=40000)
    return s.x


pA = fit(closureA, (t_att, Cs), [0.05, 7200, 0.08, 54000, 10800],
         [0, 600, 0, 0, 300], [0.5, 28800, 0.5, 144000, 72000])
pB = fit(closureB, (t_att, Cs), [0.15, 3600, 18000, 0.10, 2.0],
         [0, 300, 600, 0, 0.2], [0.6, 21600, 216000, 0.6, 20])
pC = fit(closureC, (um, Cs), [0.15, 1.0, 0.3, 0.08, 3.0],
         [0.01, 0.2, 0.05, 0.0, 0.3], [0.6, 8, 3, 0.5, 15])
for tag, Rp in (("A", closureA(pA, t_att, Cs)), ("B", closureB(pB, t_att, Cs)),
                ("C", closureC(pC, um, Cs))):
    print(f"闭合 {tag}: RMSE={float(np.sqrt(np.mean((Rp - R_att)**2))) * 1000:.4f} mm")

R_pred = closureC(pC, um, Cs)
res = R_pred - R_att
eps_c, a_c, b_c, eps_p, k_p = pC
print(f"\n采用闭合 C。标定参数（拟合附件 2，诚实标注）：ε_c={eps_c:.4f}, a={a_c:.4f}, "
      f"b={b_c:.4f}, ε_p={eps_p:.4f}, k_p={k_p:.4f}")
rmse = float(np.sqrt(np.mean(res ** 2)))
print(f"拟合优度：RMSE={rmse * 1000:.4f} mm = {rmse * 100:.5f} cm"
      f"（附件 2 量化分辨率 0.01 mm）；max|偏差|={np.max(np.abs(res)) * 1000:.4f} mm")
r0 = res - res.mean()
lag1 = float(np.sum(r0[1:] * r0[:-1]) / np.sum(r0 ** 2))
runs = int(np.sum(np.diff(np.sign(r0)) != 0)) + 1
print(f"残差结构：lag-1 自相关={lag1:.3f}，符号游程数={runs}/145（白噪参考 ~73）")

# 留出法（过拟合自检）
for name, tr, te in (("偶数点训练→奇数点留出", slice(0, None, 2), slice(1, None, 2)),
                     ("奇数点训练→偶数点留出", slice(1, None, 2), slice(0, None, 2))):
    s = fit(closureC, (um[tr], Cs[tr]), [0.15, 1.0, 0.3, 0.08, 3.0],
            [0.01, 0.2, 0.05, 0.0, 0.3], [0.6, 8, 3, 0.5, 15], target=R_att[tr])
    e_tr = float(np.sqrt(np.mean((closureC(s, um[tr], Cs[tr]) - R_att[tr]) ** 2)))
    e_te = float(np.sqrt(np.mean((closureC(s, um[te], Cs[te]) - R_att[te]) ** 2)))
    print(f"留出法[{name}]：训练 {e_tr * 1000:.4f} mm / 留出 {e_te * 1000:.4f} mm，"
          f"参数={np.round(s, 3)}")

# 结构复现核对
m_pred = R_pred / R_ideal
i_md = int(np.argmax(m_data > 1.0))
i_mp = int(np.argmax(m_pred > 1.0))
print(f"\n比率 m=R/R_ideal：data 1.000→{m_data.min():.3f}@{th[np.argmin(m_data)]:.1f}h"
      f"→{m_data[-1]:.3f}；pred 1.000→{m_pred.min():.3f}@{th[np.argmin(m_pred)]:.1f}h"
      f"→{m_pred[-1]:.3f}")
print(f"滞后交叉（m 由 <1 转 >1）：data≈{th[i_md]:.1f} h，pred≈{th[i_mp]:.1f} h；"
      f"平台：data {R_att[-1] * 100:.3f} cm，pred {R_pred[-1] * 100:.4f} cm")

# 守恒与孔隙率
V0 = 1 / RHO_SK + C0 / RHO_W
V_pred = V0 * (R_pred / R0) ** 2
V_sw = 1 / RHO_SK + um / RHO_W
V_pore = V_pred - V_sw
i_zero = int(np.argmax(V_pore > 1e-9))
print(f"\n守恒/孔隙率：V=V_s+V_w+V_pore 逐点自洽；V_pore 早期为负=基体压密（塌陷），"
      f"转正于 t≈{th[i_zero]:.1f} h，其后单调非减="
      f"{bool(np.all(np.diff(V_pore[i_zero:]) >= -1e-9))}；"
      f"末期孔隙率={V_pore[-1] / V_pred[-1] * 100:.1f}%（附件 2 隐含 ~31.0%）")

# ---------- CSV ----------
term_c, term_p = termsC(pC, um, Cs)
out = np.column_stack([t_att, R_att * 100, R_pred * 100, res * 1000,
                       R_ideal * 100, term_c * 100, term_p * 100, V_pore / V0 * 100])
np.savetxt(DATA / "endogenous_calibrated.csv", out, delimiter=",",
           header=("t_s,R_data_cm,R_pred_cm,residual_mm,R_ideal_cm,"
                   "collapse_term_cm,pore_term_cm,V_pore_frac_pct"), comments="")
print("written:", DATA / "endogenous_calibrated.csv")

# ---------- 图 ----------
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False
fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(9, 7), sharex=True,
                               gridspec_kw=dict(height_ratios=[3, 1]))
ax1.plot(th, R_att * 100, "k-", lw=2, label="附件 2（实测）")
ax1.plot(th, R_pred * 100, "r--", lw=1.8,
         label=f"内生闭合 C（5 参数标定，RMSE={rmse * 1000:.3f} mm）")
ax1.plot(th, R_ideal * 100, "b:", lw=1.4, label="失水理想收缩基线")
for x in (3.5, 21.0):
    ax1.axvline(x, color="gray", alpha=0.4, ls=":")
ax1.annotate("A 急缩段", xy=(1.0, 1.97), fontsize=9)
ax1.annotate("B 缓缩段", xy=(9, 1.45), fontsize=9)
ax1.annotate("C 平台（1.198 cm）", xy=(40, 1.21), fontsize=9)
ax1.set_ylabel("R (cm)")
ax1.legend(fontsize=9)
ax1.set_title("内生收缩闭合标定与验证（一项一机制：失水 / 毛细-膨压塌陷 / 结皮-成孔停滞）")
ax2.plot(th, res * 1000, "g-")
ax2.axhline(0, color="k", lw=0.5)
ax2.set_ylabel("残差 (mm)")
ax2.set_xlabel("t (h)")
fig.tight_layout()
SCRATCH.mkdir(exist_ok=True)
fig.savefig(SCRATCH / "内生收缩_标定验证.png", dpi=150)
print("written:", SCRATCH / "内生收缩_标定验证.png")

# ---------- (a) t_f 一致性：R_pred 替换附件 2 重解问题 4（自适应求解器） ----------
env_t, TaK, Ce = read_env()
env_hold = make_env("hold", env_t, TaK, Ce)
pchip_pred = PchipInterpolator(t_att, R_pred, extrapolate=False)
R_fun_pred = lambda tt: float(pchip_pred(min(tt, t_att[-1])))
r_a = solve_coupled4(80, env_hold, R_fun_pred, 600000.0, dt_max=15.0)
tf_a = r_a["crossing"][2] / 3600.0
print(f"\n[t_f 一致性 a] R_pred(C) 替换附件 2 重解（自适应求解器，N=80）："
      f"t_f = {tf_a:.4f} h（主解 50.8295 h，Δ = {tf_a - 50.829547:+.4f} h）", flush=True)

# ---------- (b) 耦合内生：R 由湿度场逐步预测，不读附件 2 ----------
print("[耦合内生 b] 开始（N=80，dt=60 s，一步对两半步外推）…", flush=True)
N = 80
dq = 1.0 / N
U = np.full(N, C0)
T = np.full(N, 301.15)
t = 0.0
dt_fix = 60.0
log = []


def R_cpl(U_, Cs_):
    um_ = float(np.mean(U_))
    tc, tp_ = termsC(pC, np.array([um_]), np.array([Cs_]))
    return float(R0 * np.sqrt((1 + um_) / (1 + C0)) * (1 + tc[0]) * (1 + tp_[0]))


cross_b = None
while t < 300000.0:
    Rt0 = R_cpl(U, 1.0)
    Cs_now, _, _ = surface_m4(U[-1], T[-1], env_hold(t)[1], dq, Rt0, D_of4)
    Rt = R_cpl(U, Cs_now)
    Ta_m, Ce_m = env_hold(t + dt_fix / 2)
    Ta_e, Ce_e = env_hold(t + dt_fix)
    Uf, Tf, *_ = coupled_step4(U, T, dt_fix, Ta_e, Ce_e, dq, True, Rt, PROPS4)
    Uh, Th, *_ = coupled_step4(U, T, dt_fix / 2, Ta_m, Ce_m, dq, True, Rt, PROPS4)
    Uh2, Th2, *_ = coupled_step4(Uh, Th, dt_fix / 2, Ta_e, Ce_e, dq, True, Rt, PROPS4)
    U, T = 2 * Uh2 - Uf, 2 * Th2 - Tf
    t += dt_fix
    if len(log) == 0 or t - log[-1][0] >= 1800 - 1e-9:
        log.append((t, Rt, float(np.mean(U)), float(np.max(U))))
    if float(np.max(U)) < TH:
        cross_b = t
        break
log = np.array(log)
tf_b = cross_b / 3600.0
R_b_att = np.interp(t_att, log[:, 0], log[:, 1])
rmse_b = float(np.sqrt(np.mean((R_b_att - R_att) ** 2)))
print(f"[耦合内生 b] t_f = {tf_b:.4f} h（主解 50.8295 h，Δ = {tf_b - 50.829547:+.4f} h）；"
      f"R(t) vs 附件 2：RMSE = {rmse_b * 1000:.4f} mm，末期 R = {log[-1, 1] * 100:.4f} cm")
np.savetxt(DATA / "endogenous_coupled.csv",
           np.column_stack([log[:, 0], log[:, 1] * 100, log[:, 2], log[:, 3]]),
           delimiter=",", header="t_s,R_coupled_cm,U_mean,U_max", comments="")
print("written:", DATA / "endogenous_coupled.csv")
print(f"\n[汇总] t_f：附件 2 主解 50.8295 | R_pred 重解 {tf_a:.4f} | 耦合内生 {tf_b:.4f} h")
