"""问题 1 一条龙：求解 + V1 解析对拍 + 守恒门禁 + 退化测试 + 导出 result1.xlsx。

用法：  python run_q1.py
产物：
  ../04-结果/result1.xlsx                （温度/水分浓度，1–1800 s × 0–2.0 cm，四位小数）
  ../03-数据/v1_convergence.csv          （热场 V1 对拍收敛序列）
  ../03-数据/v1_points_final.csv         （最终网格 5 位置 × 7 时刻逐点偏差）
  ../03-数据/conservation.csv            （湿分守恒门禁三档网格）
  ../03-数据/degenerate_constD.csv       （D=const 退化 vs Bessel 级数）
  ../03-数据/q1_fields.npz               （最终全场数据）
"""
from __future__ import annotations

import hashlib
import io
import sys
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

CODE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(CODE_DIR))

from solver_q1 import (ALPHA, BI_H, C0, D_of_C, HM, K, R0, T0_K,
                       apply_interp, build_interp_weights, heat_advance,
                       heat_modal, heat_surface, solve_moisture)
from analytic_heat import RobinCylinder

A_DIR = CODE_DIR.parent
ENV_XLSX = A_DIR / "01-题目" / "原始文件" / "附件1.xlsx"
TEMPLATE = A_DIR / "01-题目" / "原始文件" / "附件3" / "result1.xlsx"
OUT_XLSX = A_DIR / "04-结果" / "result1.xlsx"
DATA_DIR = A_DIR / "03-数据"

T_END = 1800
V1_GRID = [20, 40, 80, 160, 320]
N_FINAL = 320
V1_TIMES = [100, 300, 600, 900, 1200, 1500, 1800]
POS_CM = np.arange(0.0, 2.01, 0.1)          # 21 列
Q_OUT = (POS_CM * 0.01 / R0) ** 2
CHECK_TIMES = [100, 600, 1800]
CHECK_POS_M = np.array([0.0, 0.005, 0.01, 0.015, 0.02])


def read_env():
    wb = openpyxl.load_workbook(ENV_XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [(r[0], r[1], r[2]) for r in ws.iter_rows(min_row=2, values_only=True)
            if r[0] is not None and r[0] <= T_END]
    wb.close()
    t = np.array([r[0] for r in rows], dtype=float)
    Ta_C = np.array([r[1] for r in rows], dtype=float)
    Cenv = np.array([r[2] for r in rows], dtype=float)
    assert t[0] == 0 and abs(t[-1] - T_END) < 1e-9 and len(t) == 31
    # 量纲防呆：附件 1 温度是摄氏；若误加成开尔文会超界
    assert -50.0 < Ta_C.min() and Ta_C.max() < 150.0, "环境温度量纲异常（应为数十℃）"
    Ta_K = Ta_C + 273.15
    assert Ta_K.min() > 200.0, "温度必须用开尔文参与计算"
    return t, Ta_K, Cenv


def heat_series(N, env_t, env_Ta, out_times):
    """数值热场：谱分解 + 精确时间积分，返回 out_times 时刻的 (M, N) 单元中心值。"""
    lam, V, w = heat_modal(N)
    y = V.T @ (np.full(N, T0_K))
    seg = 0
    t_cur = float(env_t[0])
    out = np.zeros((len(out_times), N))

    def Ta_at(t):                                # 段内线性
        return env_Ta[seg] + (env_Ta[seg + 1] - env_Ta[seg]) * \
            (t - env_t[seg]) / (env_t[seg + 1] - env_t[seg])

    for j, tau in enumerate(out_times):
        while t_cur < tau - 1e-12:
            t_nxt = min(tau, float(env_t[seg + 1]))
            y = heat_advance(y, lam, w, t_nxt - t_cur, Ta_at(t_cur), Ta_at(t_nxt))
            t_cur = t_nxt
            if t_nxt >= float(env_t[seg + 1]) - 1e-12 and seg + 1 < len(env_t) - 1:
                seg += 1
        out[j] = V @ y
    return out


def v1_study(env_t, env_Ta):
    ana = RobinCylinder(BI_H, ALPHA, R0, n_modes=300)
    print(f"[V1] 特征展开自检: ΣA_n(中心)={ana.sum_center:.10f}, "
          f"ΣA_nJ₀(λₙ)(表面)={ana.sum_surface:.10f}")
    g = env_Ta - T0_K
    rows = []
    for N in V1_GRID:
        r_c = R0 * np.sqrt((np.arange(N) + 0.5) / N)
        T_num = heat_series(N, env_t, env_Ta, CHECK_TIMES)
        T_ana = ana.eval(r_c, CHECK_TIMES, env_t, g, T0_K)
        errs = np.max(np.abs(T_num - T_ana), axis=1)
        rows.append((N, *errs))
        print(f"[V1] N={N:4d}  单元中心逐点最大偏差(℃): "
              + "  ".join(f"t={t}:{e:.3e}" for t, e in zip(CHECK_TIMES, errs)))
    orders = []
    for k in range(len(CHECK_TIMES)):
        e = np.array([r[k + 1] for r in rows])
        ns = np.array([r[0] for r in rows])
        p = -np.polyfit(np.log(ns), np.log(e), 1)[0]
        orders.append(p)
    print("[V1] 实测收敛阶(全段拟合): "
          + "  ".join(f"t={t}:{p:.3f}" for t, p in zip(CHECK_TIMES, orders)))
    return rows, orders, ana


def v1_points_final(env_t, env_Ta, ana):
    """最终网格下 5 位置（含中心与表面重建）× 7 时刻逐点偏差。"""
    N = N_FINAL
    T_num = heat_series(N, env_t, env_Ta, V1_TIMES)          # (7, N)
    env_at = np.interp(V1_TIMES, env_t, env_Ta)
    T_surf = np.array([heat_surface(T_num[j, -1], env_at[j], N)
                       for j in range(len(V1_TIMES))])
    idx, W = build_interp_weights(N, (CHECK_POS_M / R0) ** 2)
    T_pos = apply_interp(T_num, T_surf, idx, W)              # (7, 5)
    g = env_Ta - T0_K
    T_ana = ana.eval(CHECK_POS_M, V1_TIMES, env_t, g, T0_K)
    dev = np.abs(T_pos - T_ana)
    print(f"[V1] N={N} 5 位置逐点最大偏差: {dev.max():.3e} ℃")
    return T_pos, T_ana, dev


def write_xlsx_deterministic(path, sheets):
    """sheets: {name: (header_row, data_matrix)}，四位小数写值，zip 重排去时间戳。"""
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = __import__("datetime").datetime(2000, 1, 1)
    wb.properties.creator = "run_q1"
    wb.properties.lastModifiedBy = "run_q1"
    first = True
    for name, (header, mat) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        for c, v in enumerate(header, start=1):
            ws.cell(row=1, column=c, value=v)
        for i in range(mat.shape[0]):
            for j in range(mat.shape[1]):
                cell = ws.cell(row=i + 2, column=j + 1,
                               value=round(float(mat[i, j]), 4))
                if j >= 1:
                    cell.number_format = "0.0000"
    buf = io.BytesIO()
    wb.save(buf)
    raw = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    fixed = io.BytesIO()
    with zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as z:
        for info in sorted(raw.infolist(), key=lambda x: x.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            data = raw.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl 保存时会把 modified 改写为当前时间，固定回去
                import re
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                              rb"\g<1>2000-01-01T00:00:00Z\g<2>", data)
            z.writestr(zi, data)
    path.write_bytes(fixed.getvalue())
    return hashlib.md5(fixed.getvalue()).hexdigest()


def check_against_template(path):
    tpl = openpyxl.load_workbook(TEMPLATE)
    out = openpyxl.load_workbook(path)
    ok = True
    def chk(cond, msg):
        nonlocal ok
        print(("  [OK] " if cond else "  [FAIL] ") + msg)
        ok = ok and cond
    chk(tpl.sheetnames == out.sheetnames, f"sheet 名一致: {out.sheetnames}")
    for name in tpl.sheetnames:
        wt, wo = tpl[name], out[name]
        chk(wo.cell(1, 1).value == wt.cell(1, 1).value,
            f"[{name}] A1 列表头 = {wo.cell(1, 1).value!r}")
        hdr = [wo.cell(1, c).value for c in range(2, 23)]
        expect = [round(0.1 * i, 10) for i in range(21)]
        chk(all(abs(h - e) < 1e-12 for h, e in zip(hdr, expect)),
            f"[{name}] 距离列头 0,0.1,…,2.0 共 21 列")
        chk(wo.max_row == 1801 and wo.max_column == 22,
            f"[{name}] 尺寸 {wo.max_row}×{wo.max_column}（1801×22）")
        times = [wo.cell(r, 1).value for r in range(2, 1802)]
        chk(times[0] == 1 and times[-1] == 1800 and times[1] - times[0] == 1,
            f"[{name}] A 列时间 1–1800 s 步长 1")
        bad4 = 0
        for row in wo.iter_rows(min_row=2, min_col=2, max_col=22, values_only=True):
            for v in row:
                if v is None or abs(v - round(v, 4)) > 1e-12:
                    bad4 += 1
        chk(bad4 == 0, f"[{name}] 全部数据单元格为四位小数（异常 {bad4}）")
    return ok


def main():
    DATA_DIR.mkdir(exist_ok=True)
    OUT_XLSX.parent.mkdir(exist_ok=True)
    env_t, env_Ta, env_C = read_env()
    print(f"环境数据: {len(env_t)} 点, 0–{env_t[-1]:.0f} s, "
          f"T_a {env_Ta[0]-273.15:.3f}→{env_Ta[-1]-273.15:.3f} ℃, "
          f"C_env {env_C[0]:.5f}→{env_C[-1]:.5f}")

    # ---------- V1：解析热场对拍 ----------
    conv_rows, orders, ana = v1_study(env_t, env_Ta)
    with open(DATA_DIR / "v1_convergence.csv", "w", encoding="utf-8") as f:
        f.write("N," + ",".join(f"maxerr_t{t}s" for t in CHECK_TIMES) + "\n")
        for r in conv_rows:
            f.write(",".join([str(r[0])] + [f"{e:.6e}" for e in r[1:]]) + "\n")
        f.write("order(fit)," + ",".join(f"{p:.4f}" for p in orders) + "\n")

    T_pos, T_ana_pos, dev_pos = v1_points_final(env_t, env_Ta, ana)
    with open(DATA_DIR / "v1_points_final.csv", "w", encoding="utf-8") as f:
        f.write("t_s,r_cm,T_num_C,T_analytic_C,abs_dev\n")
        for j, t in enumerate(V1_TIMES):
            for k, r in enumerate(CHECK_POS_M):
                f.write(f"{t},{r*100:.1f},{T_pos[j,k]-273.15:.8f},"
                        f"{T_ana_pos[j,k]-273.15:.8f},{dev_pos[j,k]:.3e}\n")

    # ---------- 最终热场（逐秒） ----------
    out_seconds = np.arange(1, T_END + 1, dtype=float)
    T_all = heat_series(N_FINAL, env_t, env_Ta, out_seconds)      # K, (1800, N)
    assert 200.0 < T_all.min() and T_all.max() < 400.0, "温度必须处于开尔文合理区间"
    Ta_sec = np.interp(out_seconds, env_t, env_Ta)
    Ts_sec = np.array([heat_surface(T_all[j, -1], Ta_sec[j], N_FINAL)
                       for j in range(len(out_seconds))])
    idx, W = build_interp_weights(N_FINAL, Q_OUT)
    T_cols = apply_interp(T_all, Ts_sec, idx, W) - 273.15         # ℃, (1800, 21)

    # ---------- 湿分场（逐秒，N=320） ----------
    cenv_fun = lambda t: float(np.interp(t, env_t, env_C))
    print(f"[湿分] N={N_FINAL} 自适应求解中 …")
    res = solve_moisture(N_FINAL, cenv_fun, T_END, out_times=out_seconds)
    st = res["stats"]
    print(f"[湿分] 步数 {st['n_steps']}（拒 {st['n_rej']}），"
          f"逐步守恒相对残差 max = {st['max_cons_res']:.3e}")
    C_all = np.array([res["outputs"][t] for t in out_seconds])
    Cs_sec = np.array([res["surface"][t][0] for t in out_seconds])
    C_cols = apply_interp(C_all, Cs_sec, idx, W)
    assert C_all.min() > 0 and C_all.max() <= C0 + 1e-9, "湿分越界"

    # ---------- 守恒门禁：三档网格 ----------
    cons_rows = []
    for N in (20, 80, 320):
        r_ = res if N == N_FINAL else solve_moisture(N, cenv_fun, T_END)
        s = r_["stats"]
        cons_rows.append((N, s["n_steps"], s["n_rej"], s["max_cons_res"]))
        print(f"[守恒] N={N:4d}: 逐步相对残差 max = {s['max_cons_res']:.3e}")
    with open(DATA_DIR / "conservation.csv", "w", encoding="utf-8") as f:
        f.write("N,n_steps,n_reject,max_rel_residual\n")
        for r in cons_rows:
            f.write(f"{r[0]},{r[1]},{r[2]},{r[3]:.6e}\n")

    # ---------- 退化自检：D=const 退回 Bessel 级数 ----------
    D0 = float(D_of_C(np.array(C0)))
    Bi_m = HM * R0 / D0
    ana_c = RobinCylinder(Bi_m, D0, R0, n_modes=300)
    t_test = np.array([3e4, 6e4, 1.2e5, 2.4e5])
    Cenv0 = float(env_C[0])
    Dconst = lambda C: np.full_like(np.asarray(C, dtype=float), D0)
    dDzero = lambda C: np.zeros_like(np.asarray(C, dtype=float))
    deg = solve_moisture(160, lambda t: Cenv0, t_test[-1], out_times=t_test,
                         Dfun=Dconst, dDfun=dDzero, rtol=1e-9)
    r_c160 = R0 * np.sqrt((np.arange(160) + 0.5) / 160)
    C_ana = ana_c.eval(r_c160, t_test, [0.0, t_test[-1]],
                       [Cenv0 - C0, Cenv0 - C0], C0)
    deg_errs = [float(np.max(np.abs(deg["outputs"][t] - C_ana[j])))
                for j, t in enumerate(t_test)]
    mean_num = np.mean(deg["outputs"][t_test[-1]])
    ratio = (mean_num - Cenv0) / (C0 - Cenv0)
    first = ana_c.A[0] * ana_c.mean_mode[0] * np.exp(-ana_c.mu[0] * t_test[-1])
    print(f"[退化] D=const={D0:.3e}, Bi_m={Bi_m:.3f}; 逐点最大偏差: "
          + "  ".join(f"t={t:.0f}:{e:.3e}" for t, e in zip(t_test, deg_errs)))
    print(f"[退化] t={t_test[-1]:.0f}s 平均残余比 {ratio:.6f}，首模预测 {first:.6f}，"
          f"比值 {ratio/first:.5f}")
    with open(DATA_DIR / "degenerate_constD.csv", "w", encoding="utf-8") as f:
        f.write("t_s,max_abs_dev\n")
        for t, e in zip(t_test, deg_errs):
            f.write(f"{t:.0f},{e:.6e}\n")
        f.write(f"# mean_ratio={ratio:.8f},first_mode={first:.8f},quotient={ratio/first:.6f}\n")

    # ---------- A6：中心列不得恒为常数 ----------
    print(f"[自检] 温度中心列: t=100s {T_cols[99,0]:.4f} ℃, "
          f"t=1800s {T_cols[1799,0]:.4f} ℃（解析 {T_ana_pos[0,0]-273.15:.4f} / "
          f"{T_ana_pos[-1,0]-273.15:.4f}）")
    assert T_cols[1799, 0] > T_cols[99, 0] + 0.1, "中心列异常：热场未传向中心"
    print(f"[自检] 水分中心列: t=1800s {C_cols[1799,0]:.4f} "
          f"（渗透 ~3mm，中心应基本不变）；表面列 t=1800s {C_cols[1799,-1]:.4f}")

    # ---------- 导出 result1.xlsx ----------
    time_col = np.arange(1, T_END + 1, dtype=float)
    header = ["时间\\到药材中心的距离"] + [round(0.1 * i, 10) for i in range(21)]
    mat_T = np.column_stack([time_col, T_cols])
    mat_C = np.column_stack([time_col, C_cols])
    md5 = write_xlsx_deterministic(OUT_XLSX,
                                   {"温度": (header, mat_T), "水分浓度": (header, mat_C)})
    print(f"[导出] {OUT_XLSX}  md5={md5}")
    print("[核对] 与附件3模板逐项比对：")
    ok = check_against_template(OUT_XLSX)
    assert ok

    np.savez(DATA_DIR / "q1_fields.npz",
             times=out_seconds, pos_cm=POS_CM, T_C=T_cols, C=C_cols,
             T_cell=T_all, C_cell=C_all, Ts=Ts_sec, Cs=Cs_sec,
             env_t=env_t, env_Ta=env_Ta, env_C=env_C)
    print("[完成] 全部产物已落盘。")


if __name__ == "__main__":
    main()
