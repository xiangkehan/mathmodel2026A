"""check_table6.py：逐格核对论文表 11（tables.tex 的 tab:表6）与内生交付工作簿 result4.xlsx。

用法：python check_table6.py        （在 源码 目录下运行）
退出码：0 = 全部数据格一致；1 = 存在差异（打印差异清单）。

比对口径：表 6 的行取工作簿的 6/12/…/48 h 常规行（60 s 行程的子集）与末行（t_f 数据行）；
列取 r=0/0.5/1/1.5 cm 与「药材表面」；位置超出当前表面（工作簿该格为空）在表中记 “—”。
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEX = HERE / "tables.tex"
RESULT4 = HERE.parent / "结果" / "result4.xlsx"
COLS6 = ["0", "0.5", "1", "1.5", "药材表面"]
ROW_END = "烘干结束时间"


def parse_tex_table6():
    """从 tables.tex 抽出表 6 的表头与数据行（字符串原样）。"""
    lines = TEX.read_text(encoding="utf-8").splitlines()
    i0 = next(i for i, l in enumerate(lines) if l.startswith("\\label{tab:表6}"))
    i1 = next(i for i, l in enumerate(lines[i0:], start=i0) if l.startswith("\\begin{tabular}"))
    i2 = next(i for i, l in enumerate(lines[i1:], start=i1) if l.startswith("\\end{tabular}"))
    body = [l for l in lines[i1 + 1:i2] if " & " in l]
    def clean(cell):                      # 去掉 LaTeX 换行符（\\）与首尾空白
        return cell.strip().rstrip("\\").strip()
    head = [clean(c) for c in body[0].split(" & ")]
    rows = [[clean(c) for c in l.split(" & ")] for l in body[1:]]
    return head, rows


def load_result4():
    """读内生交付工作簿 → (列名 → 列号, t → 行, 末行)。"""
    import openpyxl
    wb = openpyxl.load_workbook(RESULT4, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    body = [r for r in it if isinstance(r[0], (int, float))]
    wb.close()
    col = {str(head[j]).replace("cm", "").strip(): j for j in range(1, len(head))}
    return col, {float(r[0]): r for r in body}, body[-1]


def main():
    head, rows = parse_tex_table6()
    col, by_t, last = load_result4()
    t_last = float(last[0])
    ok, bad = [], []
    assert head[1:] == COLS6, f"表 6 列头不符：{head[1:]}"
    for cells in rows:
        label = cells[0]
        if label == ROW_END:
            src, tname = last, f"{ROW_END} (t={t_last:.3f} s)"
        else:
            th = float(label.rstrip("h"))
            src, tname = by_t.get(th * 3600.0), f"{label} (t={th * 3600:.0f} s)"
        if src is None:
            bad.append(f"{tname}: 工作簿缺该行")
            continue
        for k, cname in enumerate(COLS6, start=1):
            v = src[-1] if cname == "药材表面" else src[col[cname]]
            want = "—" if v is None else f"{float(v):.4f}"
            got = cells[k]
            if got == want:
                ok.append(f"{label:>6} × {cname:<4} = {want}")
            else:
                bad.append(f"{label:>6} × {cname:<4} 表={got}  工作簿={want}")
    print(f"[check_table6] 表 11（tab:表6）vs 内生 result4.xlsx："
          f"{len(rows)} 行 × {len(COLS6)} 列 = {len(ok) + len(bad)} 格")
    if bad:
        print(f"[FAIL] {len(bad)} 格不一致：")
        for b in bad:
            print("   ", b)
        return 1
    print(f"[PASS] 全部 {len(ok)} 个数据格一致 ✓（末行 t={t_last:.3f} s = "
          f"{t_last / 3600:.4f} h）")
    print("       抽查：", "；".join(ok[:3]), "…", ok[-1])
    return 0


if __name__ == "__main__":
    sys.exit(main())
