"""问题 2/3 一条龙：双向耦合求解 + t_f 主结果 + 全部门禁 + result2/result3。

用法：  python run_q23.py
架构（降耗）：内部自适应大步长（尾部 ~10 s），逐接受步记录 21 列重构值；
result2 的 1 s 网格与 result3 的 60 s 网格均由步记录线性插值重构（场变化慢，
插值误差 ≪ 四位小数舍入，时间误差预算另有 rtol/10 实测约束）。
主网格 N=160（T1 已证二阶收敛，N=320 仅余量）；收敛序列/外推三档用 N=40/80。

产物：
  ../04-结果/result3.xlsx        （Sheet1，60 s × 21 列，至 t_f + 追加"烘干结束时间"行）
  ../04-结果/result2.xlsx        （温度/水分浓度，1 s × 21 列，终点 = t_f + 3600 s 余量）
  ../03-数据/{tf_convergence,extrapolation,conservation23,v6_fields,v3_bound}.csv
  ../03-数据/q23_summary.json, q23_main_steps.npz
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sys
import time
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import brentq

CODE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(CODE_DIR))

from solver_q1 import R0, H, HM, T0_K, C0, build_interp_weights, apply_interp
from analytic_heat import RobinCylinder
from solver_q23 import TH, D_of, solve_coupled

A_DIR = CODE_DIR.parent
ENV_XLSX = A_DIR / "01-题目" / "原始文件" / "附件1.xlsx"
TPL_DIR = A_DIR / "01-题目" / "原始文件" / "附件3"
OUT2 = A_DIR / "04-结果" / "result2.xlsx"
OUT3 = A_DIR / "04-结果" / "result3.xlsx"
DATA_DIR = A_DIR / "03-数据"

N_MAIN = 160
MARGIN = 3600.0
POS_CM = np.arange(0.0, 2.01, 0.1)
Q_OUT = (POS_CM * 0.01 / R0) ** 2
HEADER = ["时间\\到药材中心的距离"] + [round(0.1 * i, 10) for i in range(21)]
V6_TIMES = np.array([3600.0, 36000.0, 108000.0, 180000.0])


def read_env():
    wb = openpyxl.load_workbook(ENV_XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [(r[0], r[1], r[2]) for r in ws.iter_rows(min_row=2, values_only=True)
            if r[0] is not None]
    wb.close()
    t = np.array([r[0] for r in rows], dtype=float)
    TaC = np.array([r[1] for r in rows], dtype=float)
    Ce = np.array([r[2] for r in rows], dtype=float)
    assert len(t) == 241 and t[0] == 0 and t[-1] == 14400
    assert -50.0 < TaC.min() and TaC.max() < 150.0, "附件1温度应为摄氏"
    return t, TaC + 273.15, Ce


def make_env(mode, t, TaK, Ce):
    """t>14400 s 外推三档：hold=末值保持（主）；linear=末 1 h 斜率线性外推；mean1h=末 1 h 均值。"""
    t_end = t[-1]
    i0 = int(np.searchsorted(t, t_end - 3600.0))
    if mode == "linear":
        sT = (TaK[-1] - TaK[i0]) / (t[-1] - t[i0])
        sC = (Ce[-1] - Ce[i0]) / (t[-1] - t[i0])
    elif mode == "mean1h":
        mT, mC = float(np.mean(TaK[i0:])), float(np.mean(Ce[i0:]))

    def f(tt):
        if tt <= t_end:
            return (float(np.interp(tt, t, TaK)), float(np.interp(tt, t, Ce)))
        if mode == "hold":
            return (float(TaK[-1]), float(Ce[-1]))
        if mode == "linear":
            return (float(TaK[-1] + sT * (tt - t_end)),
                    max(float(Ce[-1] + sC * (tt - t_end)), 0.0))
        return (mT, mC)
    return f


def v3_bound(t, TaK, Ce):
    """比较下界：D_max=路径点态最大（C≤C0、T≤max(T0,max Ta)），λ₁/A₁ 为
    Bi_m=h_mR0/D_max 的圆柱 Robin 第一根/中心响应系数；另给精确级数穿越与膜界。"""
    Ta_max = max(float(TaK.max()), T0_K)
    D_max = float(D_of(C0, Ta_max))
    Bi = HM * R0 / D_max
    Ce_end = float(Ce[-1])
    rc = RobinCylinder(Bi, D_max, R0, n_modes=300)
    lam1, A1 = rc.lam[0], rc.A[0]
    t_lb = R0**2 / (D_max * lam1**2) * np.log(A1 * (C0 - Ce_end) / (TH - Ce_end))
    g = [Ce_end - C0, Ce_end - C0]

    def w_center(tt):
        return float(rc.eval(np.array([0.0]), np.array([tt]), [0.0, tt + 1.0], g, C0)[0, 0])

    t_ref = brentq(lambda x: w_center(x) - TH, t_lb * 0.5, t_lb * 1.5, xtol=1e-6)
    t_film = R0 / (2 * HM) * np.log((C0 - Ce_end) / (TH - Ce_end))
    ratio2 = abs(rc.A[1]) * np.exp(-(rc.mu[1] - rc.mu[0]) * t_ref) / rc.A[0]
    return dict(D_max=D_max, Bi_m=Bi, lam1=float(lam1), A1=float(A1),
                t_lb_h=float(t_lb / 3600), t_ref_s=float(t_ref),
                t_ref_h=float(t_ref / 3600), t_film_h=float(t_film / 3600),
                Ta_max_K=Ta_max, second_mode_ratio=float(ratio2))


def run_full(N, env, rtol=1e-7, newton_below=0.6, dt_max=15.0, margin=0.0):
    """全程运行：逐接受步记录 (t, 21列U, 21列T, umax)。返回 dict。"""
    idx, W = build_interp_weights(N, Q_OUT)
    log_t, log_C, log_T, log_um = [], [], [], []

    def record(tt, U, T, Us, Ts):
        log_t.append(tt)
        log_C.append(apply_interp(U, Us, idx, W))
        log_T.append(apply_interp(T, Ts, idx, W) - 273.15)
        log_um.append(max(float(np.max(U)), 1.5 * U[0] - 0.5 * U[1]))

    res = solve_coupled(N, env, 300000.0, out_times=(), dt_max=dt_max,
                        rtolU=rtol, rtolT=rtol, newton_below=newton_below,
                        margin=margin, record=record, record_each=True)
    return dict(res=res, t=np.array(log_t),
                C=np.array(log_C), T=np.array(log_T), um=np.array(log_um))


def interp_rows(log_t, log_V, times):
    """步记录 → 指定时刻线性插值（times 升序，须在记录范围内）。"""
    times = np.asarray(times, dtype=float)
    i = np.clip(np.searchsorted(log_t, times), 1, len(log_t) - 1)
    t0, t1 = log_t[i - 1], log_t[i]
    w = np.where(t1 > t0, (times - t0) / np.maximum(t1 - t0, 1e-300), 0.0)
    return log_V[i - 1] + w[:, None] * (log_V[i] - log_V[i - 1])


def write_xlsx_stream(path, sheets):
    """sheets: [(name, header, row_iter)]；openpyxl write_only 流式 + zip 确定性重排。"""
    wb = openpyxl.Workbook(write_only=True)
    for name, header, rows in sheets:
        ws = wb.create_sheet()
        ws.title = name
        ws.append(header)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    raw = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    fixed = io.BytesIO()
    with zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as z:
        for info in sorted(raw.infolist(), key=lambda x: x.filename):
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            data = raw.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                              rb"\g<1>2000-01-01T00:00:00Z\g<2>", data)
            z.writestr(zi, data)
    path.write_bytes(fixed.getvalue())
    return hashlib.md5(fixed.getvalue()).hexdigest()


def main():
    DATA_DIR.mkdir(exist_ok=True)
    OUT2.parent.mkdir(exist_ok=True)
    t_start = time.perf_counter()
    t, TaK, Ce = read_env()
    env_hold = make_env("hold", t, TaK, Ce)
    print(f"环境: 0–14400 s 共 {len(t)} 点；末值 {TaK[-1]-273.15:.3f} ℃ / {Ce[-1]:.5f}", flush=True)

    # ---------- V3 严格下界（解析，不占求解时间） ----------
    v3 = v3_bound(t, TaK, Ce)
    print(f"[V3] D_max={v3['D_max']:.4e}, Bi_m={v3['Bi_m']:.4f}, λ₁={v3['lam1']:.6f}, "
          f"A₁={v3['A1']:.6f}")
    print(f"[V3] 首模公式下界 {v3['t_lb_h']:.4f} h；精确级数参考穿越 {v3['t_ref_h']:.4f} h；"
          f"膜界 {v3['t_film_h']:.4f} h；穿越时第二模占比 {v3['second_mode_ratio']:.2e}", flush=True)
    with open(DATA_DIR / "v3_bound.csv", "w", encoding="utf-8") as f:
        f.write("quantity,value\n")
        for k, v in v3.items():
            f.write(f"{k},{v:.10g}\n")

    # ---------- 主运行 N=160 ----------
    print(f"[主运行] N={N_MAIN} 全程 + {MARGIN:.0f}s 余量 …", flush=True)
    t0 = time.perf_counter()
    main_run = run_full(N_MAIN, env_hold, dt_max=10.0, margin=MARGIN + 20.0)
    wall_main = time.perf_counter() - t0
    cr = main_run["res"]["crossing"]
    assert cr is not None
    # 线性插值估计穿越点 + 右侧安全端点
    t_star = cr[0] + (TH - cr[1]) * (cr[2] - cr[0]) / (cr[3] - cr[1])
    t_right = cr[2]
    tf_h = t_right / 3600.0
    t_end2 = int(np.floor(t_right)) + int(MARGIN)
    st = main_run["res"]["stats"]
    print(f"[主运行] 穿越区间 [{cr[0]:.1f}, {cr[2]:.1f}] s (umax {cr[1]:.6f}→{cr[3]:.6f})，"
          f"插值 t*={t_star/3600:.6f} h，右端点 t_f={tf_h:.6f} h；"
          f"步数 {st['n_steps']}(拒{st['n_rej']})，耗时 {wall_main:.1f} s", flush=True)
    print(f"[主运行] 守恒残差 水量(逐窗max) {st['max_res_m']:.2e} / "
          f"焓(逐窗max) {st['max_res_h']:.2e} / 焓(全局吞吐) {st['glob_res_h']:.2e}；"
          f"argmax 抽样 {set(i for _, i in main_run['res']['locus'])}", flush=True)
    assert tf_h > v3["t_ref_h"], "V3 闸门未过！"
    lt, lC, lT = main_run["t"], main_run["C"], main_run["T"]

    # ---------- 舍入自洽（60 s 网格四位小数判定） ----------
    # 判定式是严格 <: 舍入值 0.1500 不算达标，首个严格 <0.15 的 60 s 行可能晚于 t_right
    t60 = np.arange(60.0, t_right + 3600.0, 60.0)
    r60 = np.round(interp_rows(lt, lC, t60), 4)
    below = np.where(np.max(r60, axis=1) < TH)[0]
    assert len(below) > 0, "余量内未找到舍入达标行"
    tf_rounded_h = float(t60[below[0]] / 3600.0)
    print(f"[舍入自洽] 60 s 四舍五入判定 {tf_rounded_h:.4f} h，"
          f"与未舍入右端点差 {abs(tf_rounded_h - tf_h):.4f} h", flush=True)

    # ---------- result3.xlsx（小，先写） ----------
    n60 = int(t_right // 60)
    r3times = np.arange(60.0, 60 * n60 + 1, 60.0)
    r3 = np.round(interp_rows(lt, lC, r3times), 4)
    r3end = np.round(interp_rows(lt, lC, [t_right])[0], 4)

    def rows3():
        for k in range(n60):
            yield [int(r3times[k])] + [float(x) for x in r3[k]]
        yield [round(float(t_right), 4)] + [float(x) for x in r3end]
    md5_3 = write_xlsx_stream(OUT3, [("Sheet1", HEADER, rows3())])
    print(f"[导出] result3.xlsx（{n60} 常规行 + 结束行 @ {t_right:.1f} s）md5={md5_3}", flush=True)

    # ---------- 收敛序列 N=40/80（N=160 用主运行） ----------
    tf_by_n = {N_MAIN: tf_h}
    st_by_n = {N_MAIN: st}
    profs = {N_MAIN: interp_rows(lt, lC, V6_TIMES)}
    for N in (40, 80):
        print(f"[收敛] N={N} …", flush=True)
        r_ = run_full(N, env_hold, dt_max=15.0)
        cr_ = r_["res"]["crossing"]
        tf_by_n[N] = cr_[2] / 3600.0
        st_by_n[N] = r_["res"]["stats"]
        profs[N] = interp_rows(r_["t"], r_["C"], V6_TIMES)
        print(f"[收敛] N={N}: t_f = {tf_by_n[N]:.6f} h, "
              f"残差 水 {st_by_n[N]['max_res_m']:.2e} 焓(全局) {st_by_n[N]['glob_res_h']:.2e}", flush=True)
    e = [tf_by_n[40], tf_by_n[80], tf_by_n[160]]
    p_num = np.log(abs(e[0] - e[1]) / abs(e[1] - e[2])) / np.log(2)
    tf_rich = e[2] + (e[2] - e[1]) / (2**p_num - 1)
    print(f"[V6] t_f 收敛阶估计 {p_num:.3f}；Richardson 外推 {tf_rich:.6f} h，"
          f"主值偏差 {abs(tf_rich - e[2]):.2e} h", flush=True)

    # ---------- 时间容差 / Picard 对照 / 外推三档（N=80） ----------
    print("[时间容差] N=80 rtol/10 …", flush=True)
    r_dt = run_full(80, env_hold, rtol=1e-8)
    tf_dt = r_dt["res"]["crossing"][2] / 3600.0
    print(f"[时间容差] t_f(1e-8) = {tf_dt:.6f} h，Δ = {abs(tf_dt - tf_by_n[80]):.2e} h", flush=True)

    print("[对照] N=80 全程 Picard …", flush=True)
    r_pc = run_full(80, env_hold, newton_below=-1.0)
    tf_pc = r_pc["res"]["crossing"][2] / 3600.0
    print(f"[对照] Picard t_f = {tf_pc:.6f} h，Δ = {abs(tf_pc - tf_by_n[80]):.2e} h", flush=True)

    ext_rows = [("hold(主)", tf_by_n[80])]
    for mode, label in (("linear", "线性外推(末1h斜率)"), ("mean1h", "末1h均值保持")):
        print(f"[外推] {label} …", flush=True)
        r_ = run_full(80, make_env(mode, t, TaK, Ce))
        ext_rows.append((label, r_["res"]["crossing"][2] / 3600.0))
    for label, tf_ in ext_rows:
        print(f"[外推] {label}: t_f = {tf_:.6f} h, Δ = {tf_ - ext_rows[0][1]:+.4f} h", flush=True)

    # ---------- V6 场量级 ----------
    v6_rows = []
    for j, tt in enumerate(V6_TIMES):
        d40 = float(np.max(np.abs(profs[40][j] - profs[160][j])))
        d80 = float(np.max(np.abs(profs[80][j] - profs[160][j])))
        v6_rows.append((tt, d40, d80))
        print(f"[V6场] t={tt:.0f}s max|ΔU_40−160|={d40:.2e} max|ΔU_80−160|={d80:.2e} "
              f"阶 {np.log(d40/d80)/np.log(2):.2f}", flush=True)

    # ---------- CSV ----------
    with open(DATA_DIR / "tf_convergence.csv", "w", encoding="utf-8") as f:
        f.write("case,t_f_h\n")
        for N in (40, 80, 160):
            f.write(f"N={N},{tf_by_n[N]:.8f}\n")
        f.write(f"N=80_rtol/10,{tf_dt:.8f}\nN=80_picard_only,{tf_pc:.8f}\n")
        f.write(f"richardson_extrap,{tf_rich:.8f}\n")
        f.write(f"interp_t_star,{t_star/3600:.8f}\nrounded_60s,{tf_rounded_h:.8f}\n")
    with open(DATA_DIR / "conservation23.csv", "w", encoding="utf-8") as f:
        f.write("N,max_rel_residual_water,max_rel_residual_enthalpy_window,glob_rel_residual_enthalpy,n_steps\n")
        for N in (40, 80, 160):
            s_ = st_by_n[N]
            f.write(f"{N},{s_['max_res_m']:.6e},{s_['max_res_h']:.6e},{s_['glob_res_h']:.6e},{s_['n_steps']}\n")
    with open(DATA_DIR / "v6_fields.csv", "w", encoding="utf-8") as f:
        f.write("t_s,maxdiff_N40_vs_N160,maxdiff_N80_vs_N160\n")
        for tt, d40, d80 in v6_rows:
            f.write(f"{tt:.0f},{d40:.6e},{d80:.6e}\n")
    with open(DATA_DIR / "extrapolation.csv", "w", encoding="utf-8") as f:
        f.write("mode,t_f_h,delta_h\n")
        for label, tf_ in ext_rows:
            f.write(f"{label},{tf_:.8f},{tf_ - ext_rows[0][1]:+.8f}\n")
    np.savez(DATA_DIR / "q23_main_steps.npz", t=lt, C=lC, T=lT, um=main_run["um"],
             crossing=np.array(cr))

    # ---------- result2.xlsx（大，write_only 流式） ----------
    print(f"[导出] result2.xlsx（{t_end2} 行 × 2 sheet，write_only 流式）…", flush=True)
    t0 = time.perf_counter()
    sec = np.arange(1.0, t_end2 + 1)
    i1 = np.clip(np.searchsorted(lt, sec), 1, len(lt) - 1)
    w1 = (sec - lt[i1 - 1]) / np.maximum(lt[i1] - lt[i1 - 1], 1e-300)
    w1 = w1[:, None]

    def rows_gen(mat):
        V0, V1 = mat[i1 - 1], mat[i1]
        for k in range(len(sec)):
            row = V0[k] + w1[k] * (V1[k] - V0[k])
            yield [int(sec[k])] + [round(float(x), 4) for x in row]
    md5_2 = write_xlsx_stream(OUT2, [
        ("温度", HEADER, rows_gen(lT)),
        ("水分浓度", HEADER, rows_gen(lC)),
    ])
    wall_w2 = time.perf_counter() - t0
    print(f"[导出] result2.xlsx md5={md5_2}，写出耗时 {wall_w2:.1f} s", flush=True)

    # ---------- 结构核对 ----------
    check_structure(t_end2, t_right, n60)

    summary = dict(tf_right_s=t_right, tf_h=tf_h, tf_star_h=t_star / 3600,
                   tf_bracket=list(cr), tf_rounded_h=tf_rounded_h,
                   tf_by_n={str(k): v for k, v in tf_by_n.items()},
                   tf_rich_h=float(tf_rich), conv_order=float(p_num),
                   tf_dt_rtol10=tf_dt, tf_picard80=tf_pc,
                   extrapolation=[(l, float(v)) for l, v in ext_rows], v3=v3,
                   wall_main_s=wall_main, wall_write2_s=wall_w2,
                   wall_total_s=time.perf_counter() - t_start,
                   md5_result2=md5_2, md5_result3=md5_3,
                   stats_main={k: float(v) for k, v in st.items() if np.isscalar(v)},
                   rows2=int(t_end2), n_main=N_MAIN)
    (DATA_DIR / "q23_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"[完成] 总耗时 {time.perf_counter() - t_start:.1f} s", flush=True)


def check_structure(t_end2, t_right, n60):
    def chk(cond, msg):
        print(("  [OK] " if cond else "  [FAIL] ") + msg, flush=True)
        assert cond, msg
    print("[核对] result3/result2 结构（对照交付口径.md §2/§3）：")
    tpl3 = openpyxl.load_workbook(TPL_DIR / "result3.xlsx")
    out3 = openpyxl.load_workbook(OUT3, read_only=True)
    chk(out3.sheetnames == tpl3.sheetnames == ["Sheet1"], f"result3 sheet: {out3.sheetnames}")
    ws = out3["Sheet1"]
    chk(ws.cell(1, 1).value == "时间\\到药材中心的距离", "result3 A1")
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, min_col=2, max_col=22))]
    chk(all(abs(h - 0.1 * i) < 1e-12 for i, h in enumerate(hdr)), "result3 21 列头 0…2.0")
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    chk(rows[0][0] == 60 and rows[1][0] == 120, "result3 A 列 60,120 起")
    chk(abs(rows[-1][0] - t_right) < 5e-4, f"result3 结束行 A={rows[-1][0]} == t_f {t_right:.4f}")
    chk(len(rows) == n60 + 1, f"result3 行数 {len(rows)} == {n60}+1")
    bad = sum(1 for r in rows for v in r[1:] if abs(v - round(v, 4)) > 1e-12)
    chk(bad == 0, f"result3 全部四位小数（异常 {bad}）")
    out3.close()
    tpl2 = openpyxl.load_workbook(TPL_DIR / "result2.xlsx")
    out2 = openpyxl.load_workbook(OUT2, read_only=True)
    chk(out2.sheetnames == tpl2.sheetnames == ["温度", "水分浓度"], f"result2 sheet: {out2.sheetnames}")
    for name in out2.sheetnames:
        ws = out2[name]
        chk(ws.cell(1, 1).value == "时间\\到药材中心的距离", f"result2[{name}] A1")
        hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, min_col=2, max_col=22))]
        chk(all(abs(h - 0.1 * i) < 1e-12 for i, h in enumerate(hdr)), f"result2[{name}] 21 列头")
        n = 0
        first3, last_a = [], None
        bad = 0
        for r in ws.iter_rows(min_row=2, values_only=True):
            n += 1
            if n <= 3:
                first3.append(r[0])
            last_a = r[0]
            if n % 997 == 0:
                bad += sum(1 for v in r[1:] if abs(v - round(v, 4)) > 1e-12)
        chk(first3 == [1, 2, 3] and last_a == t_end2,
            f"result2[{name}] A 列 1,2,3…{t_end2}（{n} 行）")
        chk(bad == 0, f"result2[{name}] 抽样四位小数（异常 {bad}）")
    out2.close()


if __name__ == "__main__":
    main()
