"""数字一致性自查（A12 结构化重写版）：期望值全部来自 03-数据 的数据文件，
main.tex / tables.tex 只作为“被核对对象”。

数据源（本模块只读，不重算）
  总账 = 03-数据/结果总账.csv             (类别, 项目, 子项) → 值 / 单位/备注；论文关键数字主源
  SPLIT = 03-数据/q4_split.csv            问题 4 效应拆分（基准＝内生几何主解）
  EVAL = 03-数据/q4_endo_eval.csv         耦合主解满 72 h 的三段误差等
  CONV = 03-数据/q4_endo_convergence.csv  问题 4 网格收敛序列、附件 2 外生版 t_f
  CAL = 03-数据/endogenous_calibrated.csv 半经验闭合标定（RMSE、平台、滞后交叉）
  LAT = 03-数据/latent_scenario.csv       潜热对照情景（# key=value 头）
  PT = 03-数据/paper_tables.csv           论文表 1–6 全部单元格
  Q23 = 03-数据/q23_summary.json          问题 3 收敛序列 / 参考时标（总账无对应键者）
  MECH = 03-数据/mechRefit.csv + mechRefit_fields.npz   §6.7 力学重标定
  RES = 04-结果/result{1,2,3,4}.xlsx      交付工作簿（md5 与行数现场核对）

两类断言
  chk(desc, literal, expect, src)  ① literal 必须出现在 main.tex+tables.tex（去 $ 后字面匹配）；
                                   ② 且与 expect（数据文件里的期望值/表达式）在“字面自身精度”
                                      上一致——即 literal 是 expect 按该位数的正确舍入
                                      （逐数比较，容差取两侧较粗的末位半个单位 ×1.01 浮点余量）。
  lit(desc, literal, src)          数据文件无对应键者（区间、复合式、外部锚点、叙述性文字）：
                                   只做字面存在性核对，注释/输出里注明出处。
失败：打印逐条差异清单并 exit 1；全过：打印统计并 exit 0。

用法：python check_numbers.py
"""
from __future__ import annotations

import csv
import hashlib
import json
import random
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
A_DIR = HERE.parent
DATA = A_DIR / "03-数据"
RESDIR = A_DIR / "04-结果"

RAW = ((HERE / "main.tex").read_text(encoding="utf-8")
       + (HERE / "tables.tex").read_text(encoding="utf-8"))
TEX = RAW.replace("$", "")          # 去数学模式符号后做字面匹配（沿用原判据）

_NUM = re.compile(r"[+-]?\d+(?:\.\d+)?(?:[eE][+-]?\d+)?")


def _latex_num(s: str) -> str:
    """把 LaTeX 科学计数写法折成 e 写法，便于取数：5.9\\times10^{-7} → 5.9e-7。"""
    s = s.replace("\\,", "").replace("\\;", "").replace("\\%", "")
    return re.sub(r"\\times\s*10\^\{?([+-]?\d+)\}?", r"e\1", s)


def _tol_of(tok: str) -> float:
    """单个数字字面量的“末位半个单位”（含 e 指数）。"""
    m = re.match(r"[+-]?\d+(?:\.(\d+))?(?:[eE]([+-]?\d+))?", tok)
    frac = len(m.group(1) or "")
    exp = int(m.group(2) or 0)
    return 0.5 * 10.0 ** (exp - frac)


def _pairs(s: str):
    """字符串中的数值与其精度：[(值, 末位半单位), …]（有序，供区间/序列逐数比较）。

    “区间”写法先归一化：1--2 / 1–2 / 1-2（数字之间的单破折号）都算区间分隔，
    不在数字前的 '-' 仍按负号处理（如 -0.286）。
    """
    t = _latex_num(s)
    t = re.sub(r"-{2,}", " ", t)                 # 双破折号 = 区间
    t = re.sub(r"(?<=\d)\s*[-–—]\s*(?=\d)", " ", t)  # 数字之间的单破折号 = 区间
    t = t.replace("–", " ").replace("—", " ")
    return [(float(m.group(0)), _tol_of(m.group(0))) for m in _NUM.finditer(t)]


# ---------------- 数据源装载 ----------------
def _load_ledger():
    d = {}
    with open(DATA / "结果总账.csv", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            d[(r["类别"].strip(), r["项目"].strip(), r["子项"].strip())] = (
                r["值"].strip(), r["单位/备注"].strip())
    return d


def _load_kv(name):
    """# key=value 头 + 可能有表体的 csv（latent_scenario.csv 形态）。"""
    out = {}
    for line in open(DATA / name, encoding="utf-8"):
        if line.startswith("# ") and "=" in line:
            k, v = line[2:].strip().split("=", 1)
            out[k] = v
        elif not line.startswith("#") and "," in line:
            head = line.strip().split(",")
            break
    return out


def _load_rows(name):
    with open(DATA / name, encoding="utf-8") as f:
        rows = [r for r in csv.reader(f) if r and not r[0].startswith("#")]
    return rows[0], rows[1:]


LEDGER = _load_ledger()
_, SPLIT_ROWS = _load_rows("q4_split.csv")
SPLIT = {r[0]: r[1] for r in SPLIT_ROWS}
_, EVAL_ROWS = _load_rows("q4_endo_eval.csv")
EVAL = {r[0]: r[1] for r in EVAL_ROWS}
_, CONV_ROWS = _load_rows("q4_endo_convergence.csv")
CONV = {r[0]: r[1] for r in CONV_ROWS}
CAL_HEAD, CAL_BODY = _load_rows("endogenous_calibrated.csv")
CAL = [dict(zip(CAL_HEAD, r)) for r in CAL_BODY]
LAT = _load_kv("latent_scenario.csv")
PT_HEAD, PT_BODY = _load_rows("paper_tables.csv")
PT = [dict(zip(PT_HEAD, r)) for r in PT_BODY]
Q23 = json.loads((DATA / "q23_summary.json").read_text(encoding="utf-8"))
MECH_HEAD, MECH_BODY = _load_rows("mechRefit.csv")
MECH = [dict(zip(MECH_HEAD, r)) for r in MECH_BODY]

FAIL: list[str] = []
N_CHK = N_LIT = 0


def _led(cat, item, sub=""):
    key = (cat, item, sub)
    if key not in LEDGER:
        raise KeyError(f"结果总账.csv 无此键：{key}（重命名/删行后须同步本脚本）")
    return LEDGER[key]


def _led_note_num(cat, item, sub=""):
    """取总账某一行的“单位/备注”里的数值（如 Δ+22.6261 h）。"""
    _, note = _led(cat, item, sub)
    return note


def _led_num(cat, item, sub=""):
    """取总账某一行的“值”列的第一个数值（值列可能带括号说明，如 2.538e-10(C≈0.108)）。"""
    v, _ = _led(cat, item, sub)
    p = _pairs(v)
    if not p:
        raise ValueError(f"总账 {cat}/{item} 的值列无数字：{v!r}")
    return p[0][0]


def chk(desc: str, literal: str, expect: str | float, src: str):
    """字面存在 + 与期望值一致（按字面精度逐数核对）。"""
    global N_CHK
    N_CHK += 1
    if literal not in TEX:
        FAIL.append(f"[缺] {desc}：正文中找不到字面 “{literal}”（期望来源：{src}）")
        return
    exp = expect if isinstance(expect, str) else repr(expect)
    if literal in exp:                              # 字面直接出现在期望串里
        return
    a, b = _pairs(literal), _pairs(exp)
    if not a or len(a) != len(b):
        FAIL.append(f"[形] {desc}：正文 “{literal}” 与 {src} 的 “{exp}” 无法逐数比对"
                    f"（数个数 {len(a)} vs {len(b)}）")
        return
    for (va, ta), (vb, tb) in zip(a, b):
        tol = max(ta, tb) * 1.01 + 1e-15
        if abs(va - vb) > tol:
            FAIL.append(f"[值] {desc}：正文 {va:g} 与 {src} 的 {vb:g} 不一致"
                        f"（差 {abs(va - vb):.3g} > 容差 {tol:.3g}）")


def lit(desc: str, literal: str, src: str):
    """字面存在性核对（数据文件无对应键；src 注明出处）。"""
    global N_LIT
    N_LIT += 1
    if literal not in TEX:
        FAIL.append(f"[缺] {desc}：正文中找不到字面 “{literal}”（来源：{src}）")


def _pt(表, 行, 列):
    for r in PT:
        if r["表"] == 表 and r["行"] == 行 and r["列"] == 列:
            return r["值"]
    raise KeyError(f"paper_tables.csv 无 {表}/{行}/{列}")


def _cal_rmse_mm():
    res = [float(r["residual_mm"]) for r in CAL]
    return (sum(x * x for x in res) / len(res)) ** 0.5


def main():
    # ============ t_f 主值族 ============
    chk("问题3 t_f 主值", "57.1799", _led("t_f", "问题3主值(未舍入右端点)"),
        "总账 t_f/问题3主值(未舍入右端点)")
    chk("问题3 t_f 插值", "57.1795", Q23["tf_star_h"], "q23_summary.json tf_star_h")
    chk("问题3 Richardson", "57.1636", _led("t_f", "问题3 Richardson外推"),
        "总账 t_f/问题3 Richardson外推")
    chk("问题3 舍入判定", "57.2333", _led("t_f", "问题3舍入判定(60s四位小数)"),
        "总账 t_f/问题3舍入判定")
    chk("问题4 t_f 主值(内生)", "50.5369", _led("t_f", "问题4主值(未舍入右端点,内生几何)")[0],
        "总账 t_f/问题4主值(内生几何)")
    chk("问题4 收敛序列(内生)", "50.5450/50.5370/50.5369",
        f"{CONV['N=40']}/{CONV['N=80']}/{CONV['N=160']}", "q4_endo_convergence.csv N=40/80/160")
    chk("问题4 交叉验证(外生版)", "50.8230", CONV["exogenous_att2_N160"],
        "q4_endo_convergence.csv exogenous_att2_N160")
    chk("内生-外生偏差", "-0.286", _led("验证", "耦合内生t_f vs 外生t_f偏差"),
        "总账 验证/耦合内生t_f vs 外生t_f偏差")

    # ============ 参考下界 ============
    chk("下界 常系数参考时标", "15.70", _led("验证", "V3问题3常系数参考时标"),
        "总账 验证/V3问题3常系数参考时标")
    chk("下界 瞬时首模参考", "19.02", _led("验证", "V3问题4瞬时首模近似参考(时间积分)"),
        "总账 验证/V3问题4瞬时首模近似参考")
    lit("下界 膜界", "11.17", "03-数据/v3_bound.csv t_film_h（总账无键）")
    lit("下界 固定 R₀ 版", "37.61", "03-数据/q4_bound.csv t_r0_h（总账无键）")
    lit("下界 R_min 常值版", "15.95", "03-数据/q4_bound.csv t_rmin_h（总账无键）")

    # ============ 验证数字（V1–V6） ============
    chk("V1 解析对拍最大偏差", "5.9\\times10^{-7}", _led("验证", "V1热场解析对拍最大偏差(N=320)"),
        "总账 验证/V1热场解析对拍最大偏差")
    chk("V1 收敛阶", "2.00/1.99/1.99", _led("验证", "V1收敛阶"), "总账 验证/V1收敛阶")
    chk("V2 尾段 ln 线性 R²", "0.9972", _led("验证", "V2尾段ln线性R2"),
        "总账 验证/V2尾段ln线性R2")
    chk("V2 尾段衰减率", "3.55\\times10^{-6}", _led("验证", "V2尾段后验诊断:衰减率γ(反演)"),
        "总账 验证/V2尾段后验诊断:衰减率γ(反演)")
    chk("V2 等效 D 反演", "2.54\\times10^{-10}", _led_num("验证", "V2尾段等效D反演(对应C)"),
        "总账 验证/V2尾段等效D反演(对应C)")
    chk("V2 反演对应 C", "0.108", _led("验证", "V2尾段等效D反演(对应C)"),
        "总账 验证/V2尾段等效D反演(对应C) 备注里的 C≈0.108")
    chk("V4 水量逐窗残差", "5.5\\times10^{-11}", _led("验证", "V4水量逐窗残差max"),
        "总账 验证/V4水量逐窗残差max")
    chk("V4 显热离散平衡残差", "1.1\\times10^{-11}",
        _led("验证", "V4冻结热容显热离散平衡全局残差"),
        "总账 验证/V4冻结热容显热离散平衡全局残差")
    chk("V5b 退化首模商", "0.99996", _led("验证", "V5b退化首模商"), "总账 验证/V5b退化首模商")
    lit("网格收敛场级残差", "3.1\\times10^{-9}", "正文 §5 表 4 网格收敛行（总账无键）")
    chk("V6 场级空间阶", "1.90--2.45", _led("验证", "V6场级空间阶"), "总账 验证/V6场级空间阶")
    chk("V5a 移动域vs固定域", "0.00", _led("验证", "V5a移动域vs固定域Δt_f"),
        "总账 验证/V5a移动域vs固定域Δt_f")

    # ============ 判据族与反验 ============
    for literal, key in (("10.6860--640.6171", "Le[附录3]"), ("31.2068--556.1413", "Le[附录4]"),
                         ("1.1805--47.7480", "Bi_m[附录3]"), ("6.3606--100.3742", "Bi_m[附录4]"),
                         ("1.0353--1.9263", "Bi_h[附录3]"), ("1.8964--3.4226", "Bi_h[附录4]")):
        chk(f"判据区间 {key}", literal, _led("判据", key), "总账 判据/" + key)
    for desc, literal, key in (
            ("判据反验 Le", "0.303", "Le反验(准稳态热场)"),
            ("判据反验 Bi_m", "46.03", "Bi_m反验(集总)"),
            ("判据反验 Bi_h", "0.075", "Bi_h反验(h→∞)")):
        chk(desc, literal, _led_note_num("判据反验", key), f"总账 判据反验/{key} 的 Δ 备注")

    # ============ 敏感性与效应拆分 ============
    chk("敏感性 C_e=0.1366", "79.8270", _led("敏感性", "C_e=0.1366(山药解吸支,题外修正)")[0],
        "总账 敏感性/C_e=0.1366")
    chk("敏感性 C_e 偏移", "22.6261", _led("敏感性", "C_e=0.1366(山药解吸支,题外修正)")[1],
        "总账 敏感性/C_e=0.1366 的 Δ 备注")
    chk("敏感性 C_e 偏移 %", "39.6", 22.626091 / 57.200904 * 100.0,
        "总账 敏感性/C_e 偏移 ÷ 末值保持主值")
    chk("环境外推 线性外推", "58.8203", _led("敏感性", "环境外推:线性外推")[0],
        "总账 敏感性/环境外推:线性外推")
    chk("环境外推 末 1h 均值", "57.5034", _led("敏感性", "环境外推:末1h均值")[0],
        "总账 敏感性/环境外推:末1h均值")
    chk("半径插值 线性", "50.8317", _led("敏感性", "半径插值线性(问题4)")[0],
        "总账 敏感性/半径插值线性(问题4)")
    chk("半径插值 两档差", "0.0022", _led("敏感性", "半径插值线性(问题4)")[1],
        "总账 敏感性/半径插值线性(问题4) 的 Δ 备注")
    chk("问题4 拆分 附4固定R", "129.1047", _led("敏感性", "问题4效应拆分:附4固定R")[0],
        "总账 敏感性/问题4效应拆分:附4固定R")
    chk("问题4 拆分 收缩效应(内生基准)", "78.57", SPLIT["shrink_effect_endo_h"],
        "q4_split.csv shrink_effect_endo_h")
    chk("问题4 拆分 收缩效应 %", "60.9", SPLIT["shrink_pct_of_fixedR_endo"],
        "q4_split.csv shrink_pct_of_fixedR_endo")
    chk("问题4 拆分 净效应 %", "11.6", SPLIT["net_pct_of_q3_endo"],
        "q4_split.csv net_pct_of_q3_endo")
    chk("问题4 拆分 净效应", "6.64", SPLIT["net_effect_endo_h"], "q4_split.csv net_effect_endo_h")

    # ============ 误差预算 ============
    chk("预算 空间网格 Q3", "0.0164", _led("误差预算", "数值误差:空间网格(Q3,N=160 Richardson偏差)")[0],
        "总账 误差预算/数值误差:空间网格(Q3,…)")
    chk("预算 时间步", "0.0014", _led("误差预算", "数值误差:时间步(rtol/10实测)")[0],
        "总账 误差预算/数值误差:时间步")
    chk("预算 输出重建", "2.6\\times10^{-4}", _led("误差预算", "数值误差:输出重建/1s插值(曲率上界)")[0],
        "总账 误差预算/数值误差:输出重建")
    chk("预算 舍入判定", "0.0701", _led("误差预算", "口径依赖:舍入判定(60s网格+四位小数)")[0],
        "总账 误差预算/口径依赖:舍入判定")
    chk("预算 环境外推", "1.6194", _led("误差预算", "环境情景:环境外推(三档最大偏移)")[0],
        "总账 误差预算/环境情景:环境外推")
    chk("预算 累计量级(问题3)", "1.71", _led("误差预算", "已检查因素累计偏移量级(问题3)")[0],
        "总账 误差预算/已检查因素累计偏移量级(问题3)")

    # ============ 叙述数字（有数据源者） ============
    chk("问题1 中心温度 1800 s", "33.5753", _pt("表1", "t=1800s", "r=0.0cm"),
        "paper_tables.csv 表1/t=1800s/r=0.0cm")
    chk("问题1 表面含水率 1800 s", "1.5102", _pt("表2", "t=1800s", "r=2.0cm"),
        "paper_tables.csv 表2/t=1800s/r=2.0cm")
    chk("材料参数组合效应倍数", "2.26",
        float(_led("敏感性", "问题4效应拆分:附4固定R")[0]) / float(_led("t_f", "问题3主值(未舍入右端点)")[0]),
        "总账 129.1047/57.1799")
    chk("环境外推 末值保持 t_f", "57.2009", _led("敏感性", "环境外推:末值保持(主)")[0],
        "总账 敏感性/环境外推:末值保持(主)")
    lit("D(C) 尾段变化倍数", "16.8", "正文 §6.4/§6.6 由附录 3 材料参数算出（总账无键）")
    lit("反演平均含水率", "0.278", "03-数据/内生收缩分析.md（总账无键）")
    lit("模型同刻平均含水率", "0.149", "03-数据/内生收缩分析.md（总账无键）")
    lit("第一问 1800 s 通量比", "5.54", "正文 §2.3 由 paper_tables.csv 表面量复算（总账无键）")

    # ---- 交付工作簿 md5 与行数（现场核对，不是誊写） ----
    for name, lit8 in (("result1.xlsx", "22eb47f7"), ("result2.xlsx", "12cde4e8"),
                       ("result3.xlsx", "fbe982ff"), ("result4.xlsx", "80c73431")):
        p = RESDIR / name
        if not p.exists():
            FAIL.append(f"[缺] {name} 不存在，无法核对 md5")
            continue
        h = hashlib.md5(p.read_bytes()).hexdigest()
        if not h.startswith(lit8):
            FAIL.append(f"[值] {name} 现 md5 {h[:8]} ≠ 正文所写 {lit8}")
        else:
            lit(f"{name} md5 锚点", lit8, f"现场计算 md5 {h[:8]}（正文 附录 B 复现清单）")
    # result4 行数（正文“3033 行”）
    try:
        import openpyxl
        wb = openpyxl.load_workbook(RESDIR / "result4.xlsx", read_only=True)
        ws = wb[wb.sheetnames[0]]
        n = sum(1 for r in ws.iter_rows(values_only=True) if r and r[0] is not None)
        wb.close()
        chk("result4 数据行数", "3033", float(n - 1), "现场读取 result4.xlsx（含表头 %d 行）" % n)
    except Exception as e:              # pragma: no cover
        FAIL.append(f"[缺] 无法读取 result4.xlsx 行数：{e}")

    # ============ 收敛序列 ============
    chk("问题3 收敛序列", "57.249/57.201/57.180",
        "/".join(f"{Q23['tf_by_n'][k]:.4f}" for k in ("40", "80", "160")),
        "q23_summary.json tf_by_n 40/80/160")

    # ============ 收缩机制分析小节（来源：03-数据/内生收缩分析.md） ============
    for s in ("0.182 cm", "9.1\\%", "0.028,0.580", "0.129", "0.155", "31\\%",
              "0.234", "1.198 cm", "61\\%", "约 16 h 交叉", "-2.4\\%", "+12\\%",
              "C_{\\mathrm{glass}}\\approx0.21"):
        lit(f"收缩机制 {s}", s, "03-数据/内生收缩分析.md（总账无键）")

    # ============ 内生收缩模型小节 ============
    for s in ("0.1108", "1.668", "1.220", "0.0709", "3.787"):
        lit(f"收缩模型参数 {s}", s, "03-数据/内生收缩模型.md §参数（总账无键）")
    chk("标定 RMSE", "0.047 mm", _cal_rmse_mm(), "endogenous_calibrated.csv 残差列现算 RMSE")
    chk("标定 max|残差|", "0.162 mm", max(abs(float(r["residual_mm"])) for r in CAL),
        "endogenous_calibrated.csv 残差列现算 max")
    chk("标定平台(72 h)", "1.1969", CAL[-1]["R_pred_cm"], "endogenous_calibrated.csv 末行 R_pred_cm")
    # 滞后交叉（m=R/R_ideal 自 <1 穿回 >1）现算；正文按“约 17 h / 约 16 h”表述
    def _cross(ratio):
        for j in range(1, len(ratio)):
            if ratio[j] > 1.0 > ratio[j - 1]:
                return j
        return None
    _rp = [float(r["R_pred_cm"]) for r in CAL]
    _ri = [float(r["R_ideal_cm"]) for r in CAL]
    _rd = [float(r["R_data_cm"]) for r in CAL]
    _j = _cross([a / b for a, b in zip(_rp, _ri)])
    _t = [float(r["t_s"]) / 3600.0 for r in CAL]
    t_x = _t[_j - 1] + (1.0 - (_rp[_j - 1] / _ri[_j - 1])) * (_t[_j] - _t[_j - 1]) / (
        (_rp[_j] / _ri[_j]) - (_rp[_j - 1] / _ri[_j - 1]))
    chk("标定滞后交叉(预测)", "约 17 h", t_x, "endogenous_calibrated.csv 现算 R_pred/R_ideal 穿越")
    chk("内生-外生相对偏差", "-0.56\\%",
        (float(SPLIT["appendix4_endogenous_R(t)"]) - float(SPLIT["appendix4_R(t)"]))
        / float(SPLIT["appendix4_R(t)"]) * 100.0, "q4_split.csv 内生与外生版之差 ÷ 外生版")
    chk("耦合主解全程 RMSE", "0.088 mm", _led("验证", "内生几何自洽R_pred vs 附件2 RMSE")[0],
        "总账 验证/内生几何自洽R_pred vs 附件2 RMSE")
    chk("耦合三段误差 0-达标", "0.065", EVAL["rmse_0_to_crossing_mm"],
        "q4_endo_eval.csv rmse_0_to_crossing_mm")
    chk("耦合三段误差 达标-72h", "0.127", EVAL["rmse_crossing_to_72h_mm"],
        "q4_endo_eval.csv rmse_crossing_to_72h_mm")
    chk("耦合主解末期半径", "1.1836", EVAL["R_pred_72h_cm"], "q4_endo_eval.csv R_pred_72h_cm")
    # 60→72 h 的末端降幅：单位易错点（cm→mm 是 ×10，不是 ×100），单独锁死
    # 该值曾被写成 0.31 mm（偏大 10 倍），源 03-数据/内生收缩模型.md 与 main.tex 两处已更正
    chk("耦合末端 60→72h 降幅", "0.031",
        (float(EVAL["R_pred_60h_cm"]) - float(EVAL["R_pred_72h_cm"])) * 10.0,
        "q4_endo_eval.csv (R_pred_60h_cm − R_pred_72h_cm) × 10 = mm")
    lit("耦合主解末期孔隙率", "30.9\\%", "03-数据/内生收缩模型.md（总账无键）")
    lit("孔隙转正时刻", "9.5 h", "03-数据/内生收缩分析.md（总账无键）")
    lit("留出法区间", "0.045--0.050 mm", "03-数据/内生收缩模型.md §留出法（总账无键）")
    for s in ("78.57", "60.9\\%", "11.6\\%"):
        lit(f"问题4 拆分复述 {s}", s, "q4_split.csv（与上文同源，正文重复出现处）")

    # ============ 参数第二路线验证小节 ============
    for s in ("1.4\\%,28.8\\%", "4.7\\%", "11.9\\%", "11.0\\%", "3.8\\%", "23.9\\%",
              "0.98\\%", "k_p", "存疑"):
        lit(f"参数验证 {s}", s, "03-数据/参数第二路线验证.md（总账无键）")

    # ============ 两条定量声明 ============
    for s in ("1.301", "R/L=8\\%", "6.25", "39", "\\rho_{sv}=\\rho/(1+C)"):
        lit(f"定量声明 {s}", s, "正文 §1.3/§2.1 几何与材料参数推导（总账无键）")

    # ============ 参考时标 / V2·V4 归位 ============
    for s in ("14.5491", "常系数参考时标", "瞬时首模近似参考"):
        lit(f"参考时标 {s}", s, "正文 §3.3（14.5491 为固定初值边界下的首模参考，总账无键）")
    for s in ("冻结热容的显热离散平衡", "后验尾段诊断"):
        lit(f"验证归位 {s}", s, "正文 §5 验证表命名（总账无键）")

    # ============ 力学升级（扩展研究）小节 ============
    res = [float(r["residual_mm"]) for r in MECH]
    hs = [float(r["t_s"]) / 3600.0 for r in MECH]
    rmse = (sum(x * x for x in res) / len(res)) ** 0.5
    seg = []
    for a, b in ((0.0, 3.5), (3.5, 21.0), (21.0, 1e9)):
        rr = [x for x, h in zip(res, hs) if a <= h < b]
        seg.append((sum(x * x for x in rr) / max(1, len(rr))) ** 0.5)
    chk("力学 RMSE", "0.522", rmse, "mechRefit.csv 残差列现算 RMSE")
    chk("力学 急缩段 RMSE", "1.193", seg[0], "mechRefit.csv 0–3.5 h")
    chk("力学 缓缩段 RMSE", "0.875", seg[1], "mechRefit.csv 3.5–21 h")
    chk("力学 平台段 RMSE", "0.163", seg[2], "mechRefit.csv >21 h")
    chk("力学 平台(模型)", "1.1747", MECH[-1]["R_pred_cm"], "mechRefit.csv 末行 R_pred_cm")
    try:
        import numpy as np
        dz = np.load(DATA / "mechRefit_fields.npz")
        beta = float(dz["fit_params"][1])
        chk("力学 β", "0.1059", beta, "mechRefit_fields.npz fit_params[1]")
        chk("力学 t_f", "51.1667", float(dz["t_f_h"]), "mechRefit_fields.npz t_f_h")
        chk("力学 g_min", "0.1446", float(np.min(dz["gS"])), "mechRefit_fields.npz gS")
        chk("力学 De", "1.0\\times10^{-3}", float(np.max(dz["DeS"])), "mechRefit_fields.npz DeS")
        jend = float(np.interp(beta, dz["terminal"][:, 0], dz["terminal"][:, 1]))
        chk("力学 末端平衡 J_end", "0.3343", jend, "mechRefit_fields.npz terminal 内插 β")
        # 正文引用的 5 个 β=0.02/0.05/0.10/0.15/0.85 对应的 J_end
        tb, tJ = dz["terminal"][:, 0], dz["terminal"][:, 1]
        pick = [float(tJ[i]) for i, b0 in enumerate(tb) if b0 in (0.02, 0.05, 0.10, 0.15, 0.85)]
        chk("力学 末端平衡表", "0.368/0.356/0.336/0.319/0.218",
            "/".join(f"{v:.3f}" for v in pick),
            "mechRefit_fields.npz terminal（β=0.02/0.05/0.10/0.15/0.85）")
    except Exception as e:              # pragma: no cover
        FAIL.append(f"[缺] 无法读取 mechRefit_fields.npz：{e}")
    for s in ("0.604", "0.212"):
        lit(f"力学 分窗误差 {s}", s, "03-数据/力学升级_修复重标定.md §4（总账无键）")
    for s in ("0.587/0.591", "0.91", "9.5\\times10^{-11}", "0.0084/0.0023/0.0013", "0.636",
              "1.717", "5.3\\times10^{-15}", "4.3\\times10^{-11}", "4.9\\times10^{-15}",
              "2.2\\times10^{-16}", "2.7\\times10^{-16}", "-0.047", "4.188", "0.042"):
        lit(f"力学 {s}", s, "03-数据/力学升级_修复重标定.md §3/§4/§6（总账无键）")
    for s in ("扩展研究", "不能声称\"力学模型复现附件 2\""):
        lit(f"力学 定性 {s}", s, "正文 §6.7 定位声明（总账无键）")

    # ============ 潜热对照情景 ============
    chk("潜热 表面相变汇 t_f", "62.5694", _led("情景(潜热A01)", "表面相变汇t_f")[0],
        "总账 情景(潜热A01)/表面相变汇t_f")
    chk("潜热 情景 Δt_f", "5.3889", _led("情景(潜热A01)", "潜热情景Δt_f")[0],
        "总账 情景(潜热A01)/潜热情景Δt_f")
    chk("潜热 逐时温度比", "29.6/29.2", _led("情景(潜热A01)", "最大逐时|ΔT_s|/|ΔT_center|")[0],
        "总账 情景(潜热A01)/最大逐时比")
    chk("潜热 最大平均含水率差", "0.322", LAT["max_dUm"], "latent_scenario.csv max_dUm")
    for s in ("声明的有效基线模型", "声明闭合"):
        lit(f"潜热 定性 {s}", s, "正文 §2.3/§6.4 口径命名（总账无键）")

    # ============ 纠错修订（正文定性表述，无数据源） ============
    for s in ("单调映射总能构造出来", "半经验收缩闭合", "插值稳健", "完整耦合标定",
              "伪影", "17.26\\%", "7.33", "9.94\\%", "约 31\\% 固体应变",
              "C_{s,\\min}=0.0488", "近平台段", "72 h 观测区间",
              "无法检验", "基线反演的等效含水率更低",
              "问题 4 主模型 = 物理启发的半经验收缩闭合内生预测半径",
              "降为验证数据", "力学驱动升级是后续正统方向"):
        lit(f"表述 {s}", s, "正文 §6.5/§6.6/§8 定性表述（总账无键）")

    # ============ paper_tables.csv 抽查 20 格（沿用固定种子） ============
    rows = [[r["表"], r["行"], r["列"], r["值"], r["单位"]] for r in PT]
    random.seed(42)
    for r in random.sample(rows, 20):
        global N_LIT
        N_LIT += 1
        if r[3] and r[3] not in TEX:
            FAIL.append(f"[缺] 论文表格缺失: {r}")

    # ============ 汇总 ============
    total = N_CHK + N_LIT
    if FAIL:
        print(f"不一致 {len(FAIL)} 处（共核对 {total} 条断言：数据源驱动 {N_CHK} + 字面/外部来源 {N_LIT}）：")
        for x in FAIL:
            print(" -", x)
        return 1
    print(f"全部通过：{total} 条断言一致（数据源驱动 {N_CHK} 条、字面/外部来源 {N_LIT} 条，"
          f"其中 tables.tex 抽查 20 格）。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
