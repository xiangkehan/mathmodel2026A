"""M4：判据族（D1）+ 敏感性矩阵（D2）+ 误差预算（D3）+ V2 尾段对拍（D4）+ 结果总账（D5）。

用法：  python run_m4.py        （需 run_q1/run_q23/run_q4 的产物 npz 已存在）
产物（../03-数据/）：
  criteria.csv       判据族表（Le/Bi_m/Bi_h 全区间 + 反验 Δt_f）
  sensitivity.csv    敏感性矩阵（每组实测 Δt_f）
  error_budget.csv   误差预算表（≥6 源 + 确定性合成）
  v2_tail.csv        V2 尾段衰减率对照
  paper_tables.csv   论文表 1–6 全部数字
  结果总账.csv        全项目关键数字唯一数据源
"""
from __future__ import annotations

import csv
import json
import sys
import time
from pathlib import Path

import numpy as np

CODE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(CODE_DIR))

from solver_q1 import R0, H, HM, T0_K, C0
from analytic_heat import RobinCylinder, robin_eigs
import solver_q23 as sq23
from solver_q23 import TH
from solver_q4 import PROPS4, solve_coupled4
from run_q23 import read_env, make_env, interp_rows
from run_q4 import read_radius

A_DIR = CODE_DIR.parent
DATA_DIR = A_DIR / "03-数据"

CE_Q3 = 0.04986          # 题目口径（末值保持）
CE_YAM = 0.1366          # 山药解吸支（题外修正，只入敏感性）
TF3_H = 57.17992518
TF4_H = 50.823025


# ======================================================================
# D1 判据族
# ======================================================================

def props3():
    return dict(rho_cp=lambda C: (650 + 128 * np.asarray(C, float)) * (1450 + 2736 * np.asarray(C, float) / (np.asarray(C, float) + 1)),
                k=lambda C: 0.21 + 0.38 * np.asarray(C, float) / (np.asarray(C, float) + 1),
                D=lambda C, T: 2.4e-3 * np.exp(-0.45 / np.asarray(C, float)) * np.exp(-3850.0 / np.asarray(T, float)))


def criteria_table():
    """Le=α/D、Bi_m=h_mR/D、Bi_h=hR/k 沿 (C,T) 路径的全区间（附3/附4）。"""
    Cg = np.linspace(0.15, 2.55, 50)
    rows = []
    for name, P in (("附录3", props3()), ("附录4", dict(rho_cp=PROPS4["rho_cp"], k=PROPS4["k_of"], D=PROPS4["D_of"]))):
        Le_all, Bim_all, Bih_all = [], [], []
        for T in (301.15, 323.315):
            a = P["k"](Cg) / P["rho_cp"](Cg)
            D = P["D"](Cg, T)
            Le_all.append(a / D)
            Bim_all.append(HM * R0 / D)
            Bih_all.append(H * R0 / P["k"](Cg))
        Le_all, Bim_all, Bih_all = map(np.concatenate, (Le_all, Bim_all, Bih_all))
        rows.append((name, Le_all.min(), Le_all.max(), Bim_all.min(), Bim_all.max(),
                     Bih_all.min(), Bih_all.max()))
    return rows


def run_patched(N, env, patch, t_end=400000.0):
    saved = {k: getattr(sq23, k) for k in patch}
    try:
        for k, v in patch.items():
            setattr(sq23, k, v)
        res = sq23.solve_coupled(N, env, t_end, out_times=(), dt_max=15.0)
    finally:
        for k, v in saved.items():
            setattr(sq23, k, v)
    return res


def run_le_simplified(N, env, t_end=450000.0):
    """Le→∞ 反验：准稳态热场 T≡T_a(t)，只求解湿分（D(C, T_a(t))，BE+Newton+对分自适应）。

    直接用 a_eff→0 的 patch 会让半离散热方程 stiff、附件 1 的 60 s 噪声折点把步长
    压死；准稳态的正确实现是温度场直接用环境值、不求解热方程。
    """
    from scipy.linalg import solve_banded
    dq = 1.0 / N
    qf = np.arange(1, N) * dq
    U = np.full(N, C0)
    t, dt = 0.0, 1e-3
    crossing = None

    def step(U, h, Ta, Ce):
        Un = U.copy()
        for _ in range(30):
            Uf = 0.5 * (Un[:-1] + Un[1:])
            Df = sq23.D_of(Uf, Ta)
            b = 4.0 * qf * Df * h / (R0**2 * dq**2)
            e = 0.5 * sq23.dD_dC(Uf, Ta) / Df * (Un[1:] - Un[:-1])
            Cs, Fs, Ds = sq23.surface_m(Un[-1], Ta, Ce, dq)
            Rs = R0**2 * dq / (8.0 * Ds) + R0 / (2.0 * HM)
            gs = h / (Rs * dq)
            fl = b * (Un[1:] - Un[:-1])
            G = Un - U - np.concatenate((fl, [gs * (Ce - Un[-1])])) + np.concatenate(([0.0], fl))
            bn = np.concatenate((b, [0.0])); en = np.concatenate((e, [0.0]))
            bp = np.concatenate(([0.0], b)); ep = np.concatenate(([0.0], e))
            J = np.zeros((3, N))
            J[0, 1:] = -bn[:-1] * (1.0 + en[:-1])
            J[1, :] = 1.0 + bn * (1.0 - en) + bp * (1.0 + ep)
            J[1, -1] += gs
            J[2, :-1] = -bn[:-1] * (1.0 - en[:-1])
            dU = solve_banded((1, 1), J, -G)
            Un = Un + dU
            if np.max(np.abs(dU)) < 1e-11:
                break
        return Un

    while t < t_end - 1e-12:
        dtc = min(dt, t_end - t)
        Ta_m, Ce_m = env(t + 0.5 * dtc)
        Ta_e, Ce_e = env(t + dtc)
        Uf_ = step(U, dtc, Ta_e, Ce_e)
        Uh = step(U, 0.5 * dtc, Ta_m, Ce_m)
        Uh2 = step(Uh, 0.5 * dtc, Ta_e, Ce_e)
        err = float(np.max(np.abs(Uh2 - Uf_)))
        tol = 1e-11 + 1e-7 * float(np.max(np.abs(U)))
        r = err / tol
        if r <= 1.0:
            umax_old = float(np.max(U))
            U = 2.0 * Uh2 - Uf_
            t += dtc
            umax = float(np.max(U))
            if crossing is None and umax < TH:
                crossing = (t - dtc, umax_old, t, umax)
                break
            dt = min(dtc * min(2.0, max(0.3, 0.9 / max(r, 1e-30) ** (1.0 / 3.0))), 15.0)
        else:
            dt = dtc * max(0.2, 0.9 / r ** (1.0 / 3.0))
    return crossing


def d1_counterproofs(env_hold):
    """三条判据的反验实测（问题 3 设定，N=80）。"""
    out = {}
    print("[D1] 反验 a：Le→∞ 准稳态热场（T≡T_a，只求湿分）…", flush=True)
    cr_le = run_le_simplified(80, env_hold)
    out["Le"] = cr_le[2] / 3600.0
    print("[D1] 反验 b：Bi_h→0 表面热阻忽略（h→1e9）…", flush=True)
    r_bh = run_patched(80, env_hold, {"H": 1e9})
    out["Bi_h"] = r_bh["crossing"][2] / 3600.0
    # 反验 c：Bi_m→0 集总（解析膜界，无需重跑）
    out["Bi_m"] = R0 / (2 * HM) * np.log((C0 - CE_Q3) / (TH - CE_Q3)) / 3600.0
    return out


# ======================================================================
# D4 / V2 尾段
# ======================================================================

def v2_tail():
    d = np.load(DATA_DIR / "q23_main_steps.npz")
    t, um, t_cross = d["t"], d["um"], float(d["crossing"][2])
    mask = (t > t_cross - 15 * 3600) & (t <= t_cross)
    y = np.log(um[mask] - CE_Q3)
    fit = np.polyfit(t[mask], y, 1)
    gamma = -float(fit[0])                           # 实测尾段衰减率 1/s
    yfit = np.polyval(fit, t[mask])
    r2 = 1.0 - float(np.sum((y - yfit) ** 2) / np.sum((y - y.mean()) ** 2))
    P3 = props3()

    def mu1(D):
        Bi = HM * R0 / D
        return D * robin_eigs(Bi, 1)[0] ** 2 / R0**2
    from scipy.optimize import brentq as bq
    D_eff = float(bq(lambda D: mu1(D) - gamma, 1e-11, 1e-8, xtol=1e-18))
    prefactor = 2.4e-3 * float(np.exp(-3850.0 / 323.315))
    C_equiv = float(-0.45 / np.log(D_eff / prefactor))   # 反演等效加权含水率
    rows = []
    for C_eff in (0.15, 0.18, 0.20, 0.25):
        D_c = float(P3["D"](C_eff, 323.315))
        Bi = HM * R0 / D_c
        lam1 = robin_eigs(Bi, 1)[0]
        rows.append((C_eff, D_c, Bi, D_c * lam1**2 / R0**2))
    return gamma, r2, D_eff, C_equiv, rows


# ======================================================================
# D3 输出重建/插值误差估计
# ======================================================================

def interp_error_estimate():
    """1 s 输出插值误差 → t_f 小时数（曲率上界法）。"""
    d = np.load(DATA_DIR / "q23_main_steps.npz")
    t, C = d["t"], d["C"]
    # 每个记录点：用左右邻步估计线性插值误差上界 |U''|·dt²/8
    err = np.zeros(len(t) - 2)
    for i in range(1, len(t) - 1):
        h0, h1 = t[i] - t[i - 1], t[i + 1] - t[i]
        if h0 <= 0 or h1 <= 0:
            continue
        d2 = (C[i + 1] - C[i]) / h1 - (C[i] - C[i - 1]) / h0
        err[i - 1] = np.max(np.abs(d2)) / (h0 + h1) * (max(h0, h1) ** 2) / 8.0 * 4.0
    max_err_C = float(np.max(err))
    # 换算 t_f：umax 在穿越处斜率
    um, cr = d["um"], d["crossing"]
    slope = abs(um[-1] - um[0]) / (t[-1] - t[0]) if len(t) > 1 else 1e-9
    i0 = int(np.searchsorted(t, cr[0])) - 2
    slope_cross = abs((um[min(i0 + 4, len(um) - 1)] - um[max(i0 - 4, 0)]) /
                      (t[min(i0 + 4, len(t) - 1)] - t[max(i0 - 4, 0)]))
    return max_err_C, max_err_C / slope_cross / 3600.0


# ======================================================================
# D5 结果总账 + 论文表
# ======================================================================

def paper_tables():
    d1 = np.load(DATA_DIR / "q1_fields.npz")
    d23 = np.load(DATA_DIR / "q23_main_steps.npz")
    d4 = np.load(DATA_DIR / "q4_main_steps.npz")
    rows = []
    pos5 = [0.0, 0.5, 1.0, 1.5, 2.0]
    # 表1/表2（问题1）：7 时刻 × 5 位置
    times1 = [100, 300, 600, 900, 1200, 1500, 1800]
    t1, T1, C1 = d1["times"], d1["T_C"], d1["C"]
    for tt in times1:
        i = int(tt) - 1
        for j, p in enumerate([0, 5, 10, 15, 20]):
            rows.append(("表1", f"t={tt}s", f"r={pos5[j]}cm", f"{T1[i, p]:.4f}", "℃"))
            rows.append(("表2", f"t={tt}s", f"r={pos5[j]}cm", f"{C1[i, p]:.4f}", "kg/kg"))
    # 表3/表4（问题2）：0.5–3.0 h × 5 位置
    lt, lC, lT = d23["t"], d23["C"], d23["T"]
    for th in (0.5, 1.0, 1.5, 2.0, 2.5, 3.0):
        vC = interp_rows(lt, lC, [th * 3600])[0]
        vT = interp_rows(lt, lT, [th * 3600])[0]
        for j, p in enumerate([0, 5, 10, 15, 20]):
            rows.append(("表3", f"t={th}h", f"r={pos5[j]}cm", f"{vT[p]:.4f}", "℃"))
            rows.append(("表4", f"t={th}h", f"r={pos5[j]}cm", f"{vC[p]:.4f}", "kg/kg"))
    # 表5（问题3）：每 6 h × 5 位置 + 结束行
    cr = d23["crossing"]
    tf_h = cr[2] / 3600.0
    for th in range(6, int(tf_h // 6) * 6 + 1, 6):
        vC = interp_rows(lt, lC, [th * 3600.0])[0]
        for j, p in enumerate([0, 5, 10, 15, 20]):
            rows.append(("表5", f"t={th}h", f"r={pos5[j]}cm", f"{vC[p]:.4f}", "kg/kg"))
    vC = interp_rows(lt, lC, [cr[2]])[0]
    for j, p in enumerate([0, 5, 10, 15, 20]):
        rows.append(("表5", "烘干结束时间", f"r={pos5[j]}cm", f"{vC[p]:.4f}", "kg/kg"))
    # 表6（问题4）：每 6 h × (0,0.5,1,1.5,药材表面) + 结束行；超出表面留空
    lt4, C4, Cs4, cr4 = d4["t"], d4["C"], d4["Cs"], d4["crossing"]
    tf4_h = cr4[2] / 3600.0
    cols6 = [(0, "0"), (5, "0.5"), (10, "1"), (15, "1.5")]
    for th in range(6, int(tf4_h // 6) * 6 + 1, 6):
        vC = interp_rows(lt4, C4, [th * 3600.0])[0]
        us = float(np.interp(th * 3600.0, lt4, Cs4))
        for p, name in cols6:
            v = "" if np.isnan(vC[p]) else f"{vC[p]:.4f}"
            rows.append(("表6", f"t={th}h", f"r={name}cm", v, "kg/kg"))
        rows.append(("表6", f"t={th}h", "药材表面", f"{us:.4f}", "kg/kg"))
    vC = interp_rows(lt4, C4, [cr4[2]])[0]
    us = float(np.interp(cr4[2], lt4, Cs4))
    for p, name in cols6:
        v = "" if np.isnan(vC[p]) else f"{vC[p]:.4f}"
        rows.append(("表6", "烘干结束时间", f"r={name}cm", v, "kg/kg"))
    rows.append(("表6", "烘干结束时间", "药材表面", f"{us:.4f}", "kg/kg"))
    return rows, tf_h, tf4_h


def _csv_kv(path):
    """读 (key,value) 两列 CSV → dict（文件不存在返回 {}）。"""
    p = Path(path)
    if not p.exists():
        return {}
    with open(p, encoding="utf-8", newline="") as f:
        rows = list(csv.reader(f))
    return {r[0]: r[1] for r in rows[1:] if len(r) >= 2}


def table6_from_result4(pt_rows):
    """把论文表 6（问题 4）的数据格替换为**内生交付工作簿** result4.xlsx 的对应格。

    result4.xlsx 结构（交付口径.md §4+§5.1）：Sheet1，A1 表头，其后 20 列 r=0.0…1.9 cm，
    末列「药材表面」；常规行为 60 s 行程（6/12/…/48 h 均为其子集），末行为 t_f 数据行。
    位置超出当前表面时单元格为空 → 该格留空（tables.tex 渲染为“—”），与交付口径一致。
    内生工作簿不存在时原样返回（保持旧行为）。
    """
    xlsx = A_DIR / "04-结果" / "result4.xlsx"
    if not xlsx.exists():
        return pt_rows
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    body = [r for r in it if isinstance(r[0], (int, float))]
    wb.close()
    col = {}
    for j in range(1, len(head)):
        col[str(head[j]).replace("cm", "")] = j
    t_last = float(body[-1][0])
    by_t = {float(r[0]): r for r in body}
    c6 = ("0", "0.5", "1", "1.5")
    out = []
    for r in pt_rows:
        if r[0] != "表6":
            out.append(r)
            continue
        tnode, cnode = r[1], r[2]
        if tnode == "烘干结束时间":
            src = body[-1]
        else:
            th = float(tnode.split("=")[1].rstrip("h"))
            src = by_t.get(th * 3600.0)
        if src is None:
            out.append(r)
            continue
        if cnode == "药材表面":
            v = src[-1]
        else:
            j = col.get(cnode.replace("r=", "").replace("cm", ""))
            v = None if j is None else src[j]
        out.append(("表6", tnode, cnode, "" if v is None else f"{float(v):.4f}", "kg/kg"))
    print(f"[D5] 论文表 6 取自内生 result4.xlsx（末行 t={t_last:.3f} s，"
          f"{len(c6) + 1} 列表头匹配）", flush=True)
    return out


def main():
    t_start = time.perf_counter()
    env_t, TaK, Ce = read_env()
    env_hold = make_env("hold", env_t, TaK, Ce)

    # ---------- D1 ----------
    print("=" * 60, "\n[D1] 判据族表 + 反验", flush=True)
    crit = criteria_table()
    cproof = d1_counterproofs(env_hold)
    tf3_80 = None  # N=80 的 t_f（反验对照基准），从 M2 记录取 57.200904
    tf3_80 = 57.20090436
    with open(DATA_DIR / "criteria.csv", "w", encoding="utf-8") as f:
        f.write("判据,物性,定义,最小值,最大值,简化决策,反验t_f_h,反验偏差_h\n")
        f.write(f"Le,附录3+4,α/D,{min(r[1] for r in crit):.4g},{max(r[2] for r in crit):.4g},"
                f"热场准稳态(T≈T_air),{cproof['Le']:.4f},{abs(cproof['Le']-tf3_80):.4f}\n")
        f.write(f"Bi_m,附录3+4,h_m·R/D,{min(r[3] for r in crit):.4g},{max(r[4] for r in crit):.4g},"
                f"集总(内扩散无阻力),{cproof['Bi_m']:.4f},{abs(cproof['Bi_m']-tf3_80):.4f}\n")
        f.write(f"Bi_h,附录3+4,h·R/k,{min(r[5] for r in crit):.4g},{max(r[6] for r in crit):.4g},"
                f"忽略表面热滞后(h→∞),{cproof['Bi_h']:.4f},{abs(cproof['Bi_h']-tf3_80):.4f}\n")
        for name, a, b, c, d_, e, g in crit:
            f.write(f"数值区间[{name}],Le,,{a:.4g},{b:.4g},,,,\n")
            f.write(f"数值区间[{name}],Bi_m,,{c:.4g},{d_:.4g},,,,\n")
            f.write(f"数值区间[{name}],Bi_h,,{e:.4g},{g:.4g},,,,\n")
    print(f"[D1] Le {crit[0][1]:.4g}–{crit[0][2]:.4g}(附3) 反验 Δt_f={abs(cproof['Le']-tf3_80):.4f} h; "
          f"Bi_m 集总 Δt_f={abs(cproof['Bi_m']-tf3_80):.4f} h; "
          f"Bi_h Δt_f={abs(cproof['Bi_h']-tf3_80):.4f} h", flush=True)

    # ---------- D2 敏感性 ----------
    print("=" * 60, "\n[D2] 敏感性矩阵", flush=True)
    sens = [("C_e=0.04986(声明闭合,主)", tf3_80, 0.0)]
    print("[D2] C_e=0.1366（山药解吸支，题外修正）…", flush=True)
    env_yam = lambda tt: (float(np.interp(tt, env_t, TaK)) if tt <= 14400 else float(TaK[-1]), CE_YAM)
    r_yam = sq23.solve_coupled(80, env_yam, 450000.0, out_times=(), dt_max=15.0)
    tf_yam = r_yam["crossing"][2] / 3600.0
    sens.append(("C_e=0.1366(山药解吸支,题外修正)", tf_yam, tf_yam - tf3_80))
    print(f"[D2] C_e=0.1366: t_f = {tf_yam:.4f} h, Δ = {tf_yam-tf3_80:+.4f} h", flush=True)

    print("[D2] 半径插值：线性（问题 4，N=80）…", flush=True)
    t_att, Rcm, R_pchip = read_radius()
    R_lin = lambda tt: float(np.interp(min(tt, t_att[-1]), t_att, Rcm * 0.01))
    r_lin = solve_coupled4(80, env_hold, R_lin, 600000.0, dt_max=15.0)
    tf_lin = r_lin["crossing"][2] / 3600.0
    tf4_80 = 50.829547
    sens.append(("半径插值PCHIP(主,问题4)", tf4_80, 0.0))
    sens.append(("半径插值线性(问题4)", tf_lin, tf_lin - tf4_80))
    print(f"[D2] 线性插值: t_f = {tf_lin:.4f} h, Δ = {tf_lin-tf4_80:+.4f} h", flush=True)
    sens += [("环境外推:末值保持(主,M2引用)", 57.20090436, 0.0),
             ("环境外推:线性外推(M2引用)", 58.82027949, 1.61937513),
             ("环境外推:末1h均值(M2引用)", 57.50341017, 0.30250581),
             ("问题4效应拆分:附4固定R(M3引用)", 129.104709, 129.104709 - tf4_80),
             ("问题4效应拆分:附4+R(t)(主,M3引用)", TF4_H, TF4_H - tf4_80),
             ("问题4效应拆分:问题3(M2引用)", TF3_H, TF3_H - tf4_80)]
    with open(DATA_DIR / "sensitivity.csv", "w", encoding="utf-8") as f:
        f.write("组别,t_f_h,delta_t_f_h\n")
        for name, tf_, d_ in sens:
            f.write(f"{name},{tf_:.6f},{d_:+.6f}\n")

    # ---------- D4 / V2 ----------
    print("=" * 60, "\n[D4] V2 尾段衰减率对照", flush=True)
    gamma, r2_fit, D_eff, C_equiv, v2rows = v2_tail()
    with open(DATA_DIR / "v2_tail.csv", "w", encoding="utf-8") as f:
        f.write("quantity,value\n")
        f.write(f"gamma_meas_1s,{gamma:.6e}\n")
        f.write(f"ln(umax-Ce)线性拟合R2,{r2_fit:.8f}\n")
        f.write(f"D_eff_反演,{D_eff:.6e}\n")
        f.write(f"C_equiv_反演,{C_equiv:.6f}\n")
        f.write("# 首模族对照: C_eff,D_eff,Bi_m,mu1_1s\n")
        for r in v2rows:
            f.write(f"# {r[0]},{r[1]:.4e},{r[2]:.4f},{r[3]:.6e}\n")
    print(f"[D4] 实测尾段衰减率 γ={gamma:.4e}/s（ln 线性 R²={r2_fit:.6f}）；"
          f"反演等效 D={D_eff:.3e} ⇔ C≈{C_equiv:.3f}（尾段剖面区间内，结构一致）", flush=True)

    # ---------- D3 误差预算 ----------
    print("=" * 60, "\n[D3] 误差预算", flush=True)
    errC, errC_h = interp_error_estimate()
    budget = [
        ("空间网格(N=160,Richardson偏差)", 0.0164, "可收敛"),
        ("时间步(rtol/10实测)", 0.0014, "可收敛"),
        ("非线性迭代(Picard/Newton对照)", 0.000001, "可收敛"),
        ("输出重建/1s插值(曲率上界)", errC_h, "可收敛"),
        ("半径插值(线性vsPCHIP,问题4)", abs(tf_lin - tf4_80), "口径依赖"),
        ("舍入(60s网格+四位小数实测)", 0.0534 + 0.0167, "固有"),
        ("环境外推(三档最大偏移)", 1.6194, "情景假设"),
    ]
    total = sum(b[1] for b in budget)
    with open(DATA_DIR / "error_budget.csv", "w", encoding="utf-8") as f:
        f.write("误差源,量级_h,类型\n")
        for name, v, typ in budget:
            f.write(f"{name},{v:.6g},{typ}\n")
        f.write(f"确定性合成(三角不等式),{total:.6g},不含C_e口径\n")
    print(f"[D3] 输出插值 {errC_h:.2e} h；合成区间 ±{total:.3f} h（不含 C_e 口径）", flush=True)

    # ---------- D5 结果总账 ----------
    print("=" * 60, "\n[D5] 结果总账 + 论文表", flush=True)
    pt_rows, tf_h, tf4_h_ = paper_tables()
    pt_rows = table6_from_result4(pt_rows)      # 论文表 6 取内生 result4.xlsx（见其 docstring）
    with open(DATA_DIR / "paper_tables.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["表", "行", "列", "值", "单位"])
        for r in pt_rows:
            w.writerow(list(r))
    q23s = json.load(open(DATA_DIR / "q23_summary.json", encoding="utf-8"))
    q4s = json.load(open(DATA_DIR / "q4_summary.json", encoding="utf-8"))
    # 问题 4 主解 = **内生几何版**：t_f / Richardson 优先取 run_q4_endo.py 的产出
    # （q4_endo_check.csv / q4_endo_convergence.csv）；无则退回外生 q4_summary.json
    tf4_main, tf4_rich = q4s["tf4_h"], q4s["tf2_rich_h"]
    tf4_tag, tf4_note = "", "h"
    _chk = _csv_kv(DATA_DIR / "q4_endo_check.csv")
    if "tf_endo_h" in _chk:
        tf4_main = float(_chk["tf_endo_h"])
        tf4_tag = "(未舍入右端点,内生几何)"
        tf4_note = f"h（附件2外生版 {q4s['tf4_h']:.6f} 作交叉验证）"
        _conv = _csv_kv(DATA_DIR / "q4_endo_convergence.csv")
        if "richardson_extrap" in _conv:
            tf4_rich = float(_conv["richardson_extrap"])
            tf4_ridx = "问题4 Richardson外推(内生几何)"
        else:
            tf4_ridx = "问题4 Richardson外推"
    else:
        tf4_ridx = "问题4 Richardson外推"
    led = []
    A = led.append
    A(("t_f", "问题3主值(未舍入右端点)", "", f"{q23s['tf_h']:.6f}", "h"))
    A(("t_f", "问题3穿越区间", "", f"[{q23s['tf_bracket'][0]:.2f},{q23s['tf_bracket'][2]:.2f}]", "s"))
    A(("t_f", "问题3 Richardson外推", "", f"{q23s['tf_rich_h']:.6f}", "h"))
    A(("t_f", "问题3舍入判定(60s四位小数)", "", f"{q23s['tf_rounded_h']:.4f}", "h"))
    A(("t_f", f"问题4主值{tf4_tag}", "", f"{tf4_main:.6f}", tf4_note))
    A(("t_f", tf4_ridx, "", f"{tf4_rich:.6f}", "h"))
    A(("t_f", "问题3外部锚点56.92/56.93/57.45偏差", "", "+0.26/+0.25/-0.27", "h"))
    for case, v in (("V1热场解析对拍最大偏差(N=320)", "5.90e-07"),
                    ("V1收敛阶", "2.000/1.993/1.990"),
                    ("V2尾段衰减率γ(实测)", f"{gamma:.4e}/s"),
                    ("V2尾段ln线性R2", f"{r2_fit:.6f}"),
                    ("V2尾段等效D反演(对应C)", f"{D_eff:.3e}(C≈{C_equiv:.3f})"),
                    ("V3问题3严格下界", "15.7021"),
                    ("V3问题4收缩下界(时间积分)", "19.016"),
                    ("V4水量逐窗残差max", "5.54e-11"),
                    ("V4总焓全局吞吐残差", "1.05e-11"),
                    ("V5a移动域vs固定域Δt_f", "0.00"),
                    ("V5b退化首模商", "0.99996"),
                    ("V6场级空间阶", "1.90-2.45")):
        A(("验证", case, "", v, "见报告"))
    for name, tf_, d_ in sens:
        A(("敏感性", name, "", f"{tf_:.4f}", f"Δ{d_:+.4f} h"))
    for name, v, typ in budget:
        A(("误差预算", name, "", f"{v:.4g}", typ))
    A(("误差预算", "确定性合成区间", "", f"±{total:.4f}", "h(不含C_e口径)"))
    # 潜热对照情景（审计 A01；02-代码/latent_scenario.py 产出，# key=value 头；缺失则跳过）
    _lat = {}
    _lp = DATA_DIR / "latent_scenario.csv"
    if _lp.exists():
        for line in open(_lp, encoding="utf-8"):
            if line.startswith("# ") and "=" in line:
                k, v = line[2:].strip().split("=", 1)
                _lat[k] = v
            elif not line.startswith("#"):
                break
    if "t_f_latent_h" in _lat:
        A(("情景(潜热A01)", "无潜热基线t_f(N=160)", "", f"{float(_lat['t_f_base_h']):.4f}", "h(基线断言57.1799)"))
        A(("情景(潜热A01)", "表面相变汇t_f", "", f"{float(_lat['t_f_latent_h']):.4f}",
           "h(情景假设:L_v=2.4e6,ρ_d,s表面局部)"))
        A(("情景(潜热A01)", "潜热情景Δt_f", "", f"{float(_lat['delta_t_f_h']):+.4f}", "h(不进主值)"))
        A(("情景(潜热A01)", "最大逐时|ΔT_s|/|ΔT_center|", "",
           f"{float(_lat['max_dTs_C']):.2f}/{float(_lat['max_dTcenter_C']):.2f}", "℃"))
    for name, a2, b2, c2, d2_, e2, g2 in crit:
        A(("判据", f"Le[{name}]", "", f"{a2:.4g}–{b2:.4g}", "α/D"))
        A(("判据", f"Bi_m[{name}]", "", f"{c2:.4g}–{d2_:.4g}", "h_m·R/D"))
        A(("判据", f"Bi_h[{name}]", "", f"{e2:.4g}–{g2:.4g}", "h·R/k"))
    for case, v in (("Le反验(准稳态热场)", cproof["Le"]), ("Bi_m反验(集总)", cproof["Bi_m"]),
                    ("Bi_h反验(h→∞)", cproof["Bi_h"])):
        A(("判据反验", case, "", f"{v:.4f}", f"Δ{abs(v-tf3_80):+.4f} h"))
    for r in pt_rows:
        A(("论文表", f"{r[0]}[{r[1]}]", r[2], r[3], r[4]))
    with open(DATA_DIR / "结果总账.csv", "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)                       # csv 引号封装：含逗号的区间单元格不再被拆列
        w.writerow(["类别", "项目", "子项", "值", "单位/备注"])
        for r in led:
            w.writerow([str(x) for x in r])
    print(f"[D5] 结果总账 {len(led)} 行；论文表 {len(pt_rows)} 格", flush=True)
    print(f"[完成] 总耗时 {time.perf_counter() - t_start:.1f} s", flush=True)


if __name__ == "__main__":
    main()
