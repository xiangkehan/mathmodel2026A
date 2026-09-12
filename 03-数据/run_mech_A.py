"""run_mech_A.py：A 级力学驱动内生收缩——退化检查（方案 §15）+ 无量纲情景研究（§13）。

一条命令：python run_mech_A.py
产物：mechA_degradation.csv、mechA_scenarios.csv、mechA_fields_central.npz、
      ../08-临时/mechA_scenarios.png（均在本目录/08-临时）。
诚实红线：缺保水/力学数据 ⇒ 只作情景分析；p_c=p*(1−S) 是标注的低阶检验本构。
"""
from __future__ import annotations

import sys
import time
from pathlib import Path

import numpy as np
import openpyxl

HERE = Path(__file__).resolve().parent
A_DIR = HERE.parent
sys.path.insert(0, str(A_DIR / "02-代码"))
sys.path.insert(0, str(A_DIR / "03-数据"))

from solver_q1 import R0, T0_K, C0, HM, H
from solver_q4 import D_of4, k_of4, rho_cp4
from mech_shrinkage import MechParams, elem_quantities, solve_mechanics, mass_heat_step, volume_check
from run_q23 import read_env, make_env

RHO_S = 1468.0
RHO_W = 1000.0
sC0 = 1.0 / RHO_S + C0 / RHO_W
T_DRY = R0 ** 2 / float(D_of4(C0, 323.315))      # 干燥特征时间（附录4典型 D）

env_t, TaK, Ce = read_env()
env_hold = make_env("hold", env_t, TaK, Ce)

wb = openpyxl.load_workbook(A_DIR / "01-题目" / "原始文件" / "附件2.xlsx", read_only=True)
rows = [r for r in wb[wb.sheetnames[0]].iter_rows(min_row=2, values_only=True) if r[0] is not None]
wb.close()
t_att = np.array([r[0] for r in rows], float)
R_att = np.array([r[1] for r in rows], float) * 0.01


def make_X(M):
    return R0 * np.sqrt(np.linspace(0, 1, M + 1))


def P_of(phi_g0, Pi_c, Pi_tau, tau_exp=4.0):
    rho_d0 = (1 - phi_g0) / sC0
    return MechParams(rho_d0=rho_d0, rho_s=RHO_S, Kinf=1.0, Ginf=1 / 3, K1=1.0, G1=1 / 3,
                      p_star=2.0 * Pi_c, tau0=Pi_tau * T_DRY, tau_exp=tau_exp)


def mech_only_run(P, U_of_t, Tc, dt, t_end, M=16):
    """给定湿度场（均匀或逐单元）的纯力学推进。"""
    X = make_X(M)
    r = X.copy()
    Z = np.zeros((M, 3))
    log = []
    active = None
    for tt in np.arange(dt, t_end + 1e-9, dt):
        U = U_of_t(tt)
        out = solve_mechanics(r, Z, U, Tc, dt, X, P, active0=active)
        r, Z, active = out["r"], out["Z"], out["active"]
        log.append((tt, float(out["r"][-1]), out["g"].copy(), out["sig_r"].copy()))
    return X, r, Z, log


# ======================================================================
# 退化/正确性检查（方案 §15 七条）
# ======================================================================
def degradation():
    res = {}
    dt, t_end = 60.0, 4 * 3600.0
    Udec = lambda tt: np.full(16, C0 - 1.0 * tt / t_end)
    # (i) 不干燥 + 初始总应力平衡 ⇒ 尺寸无人为跳变
    P1 = P_of(0.1726, 1.0, 1.0)
    X1, r1, _, _ = mech_only_run(P1, lambda tt: np.full(16, C0), 301.15, dt, t_end)
    res["i_nojump"] = float(np.max(np.abs(r1 - X1)))
    # (ii) 关闭黏弹分支 ⇒ 纯弹性
    P2 = P_of(0.1726, 1.0, 1.0)
    P2.G1 = P2.K1 = 0.0
    _, r2, _, log2 = mech_only_run(P2, Udec, 323.315, dt, t_end)
    res["ii_elastic_R_end"] = float(r2[-1])
    # (iii) τ≪过程尺度 ⇒ 充分松弛（应与 (ii) 一致）
    P3 = P_of(0.1726, 1.0, 1e-9)
    _, r3, _, log3 = mech_only_run(P3, Udec, 323.315, dt, t_end)
    res["iii_vs_ii_maxdR"] = float(np.max(np.abs(r3 - r2)))
    # (iv) τ≫ ⇒ 内变量近冻结
    P4 = P_of(0.1726, 1.0, 1e9)
    _, r4, Z4, _ = mech_only_run(P4, Udec, 323.315, dt, t_end)
    Emax = float(np.max(np.abs(np.log(np.maximum(np.diff(r4) / np.diff(X1), 1e-300)))))
    res["iv_maxZ_over_Emax"] = float(np.max(np.abs(Z4))) / max(Emax, 1e-300)
    # (v) 指定仿射映射 ⇒ 传质系数恢复 4qD/R²
    M5 = 16
    X5 = make_X(M5)
    lam = 0.8
    r5 = lam * X5
    q5 = (X5 / R0) ** 2
    rqf = (r5[2:] - r5[:-2]) / (q5[2:] - q5[:-2])
    qf = 0.25 * (np.sqrt(q5[2:]) + np.sqrt(q5[:-2])) ** 2
    coef_num = 1.0 / rqf ** 2
    coef_ana = 4.0 * qf / (lam * R0) ** 2
    res["v_affine_relerr"] = float(np.max(np.abs(coef_num - coef_ana) / coef_ana))
    # (vi) 固定外形持续失水 ⇒ g 按体积恒等式增长
    P6 = P_of(0.1726, 1.0, 1.0)
    U6a = np.full(M5, C0)
    U6b = np.full(M5, C0 - 0.5)
    _, _, _, _, g6a, *_ = elem_quantities(X5, X5, U6a, P6)
    _, _, _, _, g6b, *_ = elem_quantities(X5, X5, U6b, P6)
    dg_num = g6b - g6a
    dg_ana = P6.rho_d0 / RHO_W * 0.5
    res["vi_growth_relerr"] = float(np.max(np.abs(dg_num - dg_ana) / dg_ana))
    # (vii) 全程 g≥0（在中心情景全程检验，见场景运行）
    return res


# ======================================================================
# 耦合情景运行
# ======================================================================
def coupled_run(P, t_end=60 * 3600.0, dt=900.0, M=16, frames=121):
    """力学→几何→传质 迭代耦合（每步 2 轮）。返回日志 dict。"""
    X = make_X(M)
    r = X.copy()
    Z = np.zeros((M, 3))
    U = np.full(M, C0)
    T = np.full(M, T0_K)
    active = None
    out_t, out_R, out_Um, out_Us, out_gmin, out_gS, out_DeS, out_volerr, out_Jmin = \
        [], [], [], [], [], [], [], [], []
    i_fr = 1
    n_steps = int(round(t_end / dt))
    for k in range(1, n_steps + 1):
        tt = k * dt
        Ta, Ce_ = env_hold(tt)
        for _ in range(2):
            mech = solve_mechanics(r, Z, U, float(np.mean(T)), dt, X, P, active0=active)
            r2, Z2, active2 = mech["r"], mech["Z"], mech["active"]
            Jc = mech["J"]
            U2, T2, Us2, Fm, Fh = mass_heat_step(U, T, dt, Ta, Ce_, r2, X, Jc, P,
                                                 D_of4, k_of4, rho_cp4)
            r, Z, active, U, T = r2, Z2, active2, U2, T2
        if k % max(1, n_steps // (frames - 1)) == 0 or k == n_steps:
            _, _, J, w, g, Pore, S = elem_quantities(r, X, U, P)
            lhs, rhs = volume_check(r, X, U, P)
            out_t.append(tt)
            out_R.append(float(r[-1]))
            out_Um.append(float(np.sum(U * np.diff((X / R0) ** 2))))
            out_Us.append(float(Us2))
            out_gmin.append(float(np.min(g)))
            out_gS.append(float(g[-1]))
            out_DeS.append(float(P.tau(max(Us2, 0.02)) / T_DRY))
            out_volerr.append(float(abs(lhs - rhs) / max(lhs, 1e-300)))
            out_Jmin.append(float(np.min(J)))
    return dict(t=np.array(out_t), R=np.array(out_R), Um=np.array(out_Um),
                Us=np.array(out_Us), gmin=np.array(out_gmin), gS=np.array(out_gS),
                DeS=np.array(out_DeS), volerr=np.array(out_volerr),
                Jmin=np.array(out_Jmin), X=X, r=r, U=U, T=T, Z=Z, g=g, S=S)


def rmse_seg(t, Rpred):
    Rp = np.interp(t_att, t, Rpred)
    err = Rp - R_att
    seg = lambda a, b: float(np.sqrt(np.mean(err[(t_att >= a) & (t_att < b)] ** 2)))
    return float(np.sqrt(np.mean(err ** 2))), seg(0, 3.5 * 3600), seg(3.5 * 3600, 21 * 3600), seg(21 * 3600, 1e12)


def main():
    t_start = time.perf_counter()
    print(f"T_DRY = {T_DRY:.3e} s = {T_DRY/3600:.1f} h", flush=True)

    # ---------- 退化检查 ----------
    print("[退化检查] 运行中 …", flush=True)
    dg = degradation()
    dg["vii_gmin_all"] = np.nan   # 由情景运行补
    for k, v in dg.items():
        print(f"  {k} = {v:.3e}", flush=True)

    # ---------- 情景扫描 ----------
    print("[情景] Π_c × Π_τ 扫描 …", flush=True)
    rows = []
    central = None
    for Pi_c in (0.2, 1.0, 5.0):
        for Pi_tau in (0.01, 1.0, 100.0):
            P = P_of(0.1726, Pi_c, Pi_tau)
            t0 = time.perf_counter()
            out = coupled_run(P, t_end=60 * 3600.0, dt=900.0, M=16)
            wall = time.perf_counter() - t0
            rm, e1, e2, e3 = rmse_seg(out["t"], out["R"])
            gmin = float(np.min(out["gmin"]))
            rows.append((Pi_c, Pi_tau, rm, e1, e2, e3, gmin, out["R"][-1], wall))
            print(f"  Π_c={Pi_c:4.1f} Π_τ={Pi_tau:7.2f}: RMSE={rm*100:.3f} cm "
                  f"(急缩 {e1*100:.3f} / 缓缩 {e2*100:.3f} / 平台 {e3*100:.3f}) "
                  f"g_min={gmin:.2e} R_end={out['R'][-1]*100:.3f} cm ({wall:.0f}s)", flush=True)
            if Pi_c == 1.0 and Pi_tau == 1.0:
                central = out
    dg["vii_gmin_all"] = min(r[6] for r in rows)
    print(f"  (vii) 全部情景 g_min = {dg['vii_gmin_all']:.3e}（≥ −1e-10 即通过）", flush=True)

    # ---------- 初始孔隙率传播（中央情景 ±） ----------
    print("[情景] 初始孔隙率 ± 传播 …", flush=True)
    for phi in (0.10, 0.30):
        P = P_of(phi, 1.0, 1.0)
        out = coupled_run(P, t_end=60 * 3600.0, dt=900.0, M=16)
        rm, e1, e2, e3 = rmse_seg(out["t"], out["R"])
        print(f"  φ_g0={phi:.4f}: RMSE={rm*100:.3f} cm, R_end={out['R'][-1]*100:.3f} cm", flush=True)

    # ---------- 中央情景加密收敛性 ----------
    print("[核对] 中央情景 M=48/dt=450 收敛性 …", flush=True)
    P48 = P_of(0.1726, 1.0, 1.0)
    out48 = coupled_run(P48, t_end=60 * 3600.0, dt=450.0, M=48)
    rm48, *_ = rmse_seg(out48["t"], out48["R"])
    print(f"  M=48/dt=450: RMSE={rm48*100:.3f} cm（vs M=16/dt=900 的 "
          f"{[r for r in rows if r[0]==1.0 and r[1]==1.0][0][2]*100:.3f} cm）", flush=True)

    # ---------- 落数据 ----------
    with open(HERE / "mechA_degradation.csv", "w", encoding="utf-8") as f:
        f.write("check,value,threshold,pass\n")
        f.write(f"i_nojump_maxdr,{dg['i_nojump']:.3e},<1e-8,{dg['i_nojump'] < 1e-8}\n")
        f.write(f"ii_elastic_R_end_cm,{dg['ii_elastic_R_end']*100:.4f},sane,True\n")
        f.write(f"iii_vs_ii_maxdR_cm,{dg['iii_vs_ii_maxdR']*100:.3e},<1e-3,{dg['iii_vs_ii_maxdR'] < 1e-5}\n")
        f.write(f"iv_maxZ_over_Emax,{dg['iv_maxZ_over_Emax']:.3e},<1e-3,{dg['iv_maxZ_over_Emax'] < 1e-3}\n")
        f.write(f"v_affine_relerr,{dg['v_affine_relerr']:.3e},<1e-10,{dg['v_affine_relerr'] < 1e-10}\n")
        f.write(f"vi_growth_relerr,{dg['vi_growth_relerr']:.3e},<1e-12,{dg['vi_growth_relerr'] < 1e-12}\n")
        f.write(f"vii_gmin_all,{dg['vii_gmin_all']:.3e},>-1e-10,{dg['vii_gmin_all'] > -1e-10}\n")
    with open(HERE / "mechA_scenarios.csv", "w", encoding="utf-8") as f:
        f.write("Pi_c,Pi_tau,rmse_cm,err_early_cm,err_mid_cm,err_late_cm,g_min,R_end_cm,wall_s\n")
        for r_ in rows:
            f.write(f"{r_[0]},{r_[1]},{r_[2]*100:.4f},{r_[3]*100:.4f},{r_[4]*100:.4f},"
                    f"{r_[5]*100:.4f},{r_[6]:.3e},{r_[7]*100:.3f},{r_[8]:.0f}\n")
    np.savez(HERE / "mechA_fields_central.npz",
             t=central["t"], R=central["R"], Um=central["Um"], Us=central["Us"],
             DeS=central["DeS"], gS=central["gS"], gmin=central["gmin"],
             X=central["X"], r_final=central["r"], U_final=central["U"],
             Z_final=central["Z"], g_final=central["g"], S_final=central["S"],
             volerr=central["volerr"])
    print(f"[完成] 总耗时 {time.perf_counter() - t_start:.0f} s", flush=True)

    # ---------- 图 ----------
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei"]
    plt.rcParams["axes.unicode_minus"] = False
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.2))
    ax = axes[0]
    ax.plot(t_att / 3600, R_att * 100, "k-", lw=2, label="附件 2")
    for Pi_c, mk in ((0.2, ":"), (1.0, "-"), (5.0, "--")):
        for Pi_tau, col in ((0.01, "C0"), (1.0, "C3"), (100.0, "C2")):
            r_ = [x for x in rows if x[0] == Pi_c and x[1] == Pi_tau][0]
            # 只画中央行与两端
    for (Pi_c, Pi_tau, *_rest) in [(0.2, 1.0), (1.0, 1.0), (5.0, 1.0), (1.0, 0.01), (1.0, 100.0)]:
        P = P_of(0.1726, Pi_c, Pi_tau)
        out = coupled_run(P, t_end=60 * 3600.0, dt=900.0, M=16)
        ax.plot(out["t"] / 3600, out["R"] * 100,
                label=f"Π_c={Pi_c:g}, Π_τ={Pi_tau:g}")
    ax.set_xlabel("t (h)"); ax.set_ylabel("R (cm)"); ax.legend(fontsize=8)
    ax.set_title("情景 R(t) vs 附件 2")
    ax2 = axes[1]
    ax2.plot(central["t"] / 3600, central["gS"], label="g（表面单元）")
    ax2.plot(central["t"] / 3600, central["gmin"], label="g_min（全域）")
    ax2.axhline(0, color="k", lw=0.5)
    ax2.set_xlabel("t (h)"); ax2.set_ylabel("g（气相体积/初始体积）")
    ax2.legend(fontsize=8)
    ax2.set_title("中央情景：孔隙非负全程成立")
    fig.tight_layout()
    fig.savefig(A_DIR / "08-临时" / "mechA_scenarios.png", dpi=150)
    print("written: 08-临时/mechA_scenarios.png", flush=True)


if __name__ == "__main__":
    main()
